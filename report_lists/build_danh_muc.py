from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parent / "Danh_muc_hinh_ve_va_bang_bieu.docx"

FIGURES = [
    ("Hình 2.1.1. Sơ đồ Use Cases Buyer", 17),
    ("Hình 2.1.2. Sơ đồ Use Cases Seller", 18),
    ("Hình 2.1.3. Sơ đồ Use Cases Admin", 18),
    ("Hình 2.2.1. Sơ đồ quản lí tài khoản", 19),
    ("Hình 2.2.2. Sơ đồ tạo build/request", 20),
    ("Hình 2.2.3. Sơ đồ Seller xử lý request", 21),
    ("Hình 2.2.4. Sơ đồ Admin quản trị", 22),
    ("Hình 2.2.5. Sơ đồ Chat", 23),
    ("Hình 2.2.6. Sơ đồ Activity", 24),
    ("Hình 2.3.1. Định tuyến tài khoản và vai trò", 25),
    ("Hình 2.3.2. Người mua tạo build và gửi yêu cầu", 25),
    ("Hình 2.3.3. Quy trình của người bán và tóm tắt QC", 26),
    ("Hình 2.3.4. Quản lý quản trị", 27),
    ("Hình 2.3.5. Phê duyệt đơn đăng ký của người mua thành người bán", 28),
    ("Hình 2.3.6. Trò chuyện giữa những người dùng", 29),
    ("Hình 2.4. Sơ đồ ERD tổng quát của hệ thống Custom Keyboard Builder", 30),
    ("Hình 2.5. Sơ đồ ERD chi tiết gồm 21 bảng và 33 khóa ngoại", 31),
    ("Hình 3.1. Giao diện đăng nhập của ứng dụng", 42),
    ("Hình 3.2. Giao diện đăng ký tài khoản Buyer", 42),
    ("Hình 3.3. Giao diện danh sách Build và trạng thái request", 43),
    ("Hình 3.4. Giao diện Buyer Dashboard và khu vực cấu hình build", 43),
    ("Hình 3.5. Giao diện danh sách Request đã gửi và theo dõi trạng thái", 44),
    ("Hình 3.6. Giao diện màn hình Buyer chat với Seller", 44),
    ("Hình 3.7. Giao diện đăng ký trở thành Seller", 45),
    ("Hình 3.8. Giao diện Seller Dashboard và xử lý yêu cầu lắp ráp", 45),
    ("Hình 3.9. Giao diện kiểm tra QC bàn phím", 46),
    ("Hình 3.10. Giao diện phân tích doanh thu và số đơn của Seller", 46),
    ("Hình 3.11. Giao diện Admin Dashboard và quản trị hệ thống", 47),
    ("Hình 3.12. Giao diện quản lý người dùng của Admin", 47),
    ("Hình 3.13. Giao diện quản lý Seller của Admin", 48),
    ("Hình 3.14. Giao diện xử lý đơn xin làm Seller của Admin", 48),
    ("Hình 3.15. Giao diện quản lý Brand bàn phím", 49),
    ("Hình 3.16. Giao diện quản lý linh kiện cho từng build", 49),
    ("Hình 3.17. Giao diện quản lý Audit log người dùng", 50),
    ("Hình 3.18. Giao diện Admin chat với Seller", 50),
]

TABLES = [
    ("Bảng 0.1. Phân công nhiệm vụ nhóm", 7),
    ("Bảng 1.1. Yêu cầu về dữ liệu", 9),
    ("Bảng 1.2. Tác nhân và mục tiêu sử dụng hệ thống", 10),
    ("Bảng 1.3. Yêu cầu chức năng chính", 10),
    ("Bảng 1.4. Yêu cầu phi chức năng", 11),
    ("Bảng 2.1. Nhóm bảng trong cơ sở dữ liệu", 30),
    ("Bảng 3.1. Thứ tự triển khai 21 bảng theo nhóm phụ thuộc", 32),
    ("Bảng 3.2. Đối chiếu giao diện với mã nguồn", 40),
    ("Bảng 3.4. Đánh giá mức độ hoàn thành theo mục tiêu", 51),
]


def set_font(run, name, size_pt, bold=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def set_keep_lines(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep_lines = ppr.find(qn("w:keepLines"))
    if keep_lines is None:
        keep_lines = OxmlElement("w:keepLines")
        ppr.append(keep_lines)
    widow = ppr.find(qn("w:widowControl"))
    if widow is None:
        widow = OxmlElement("w:widowControl")
        ppr.append(widow)
    widow.set(qn("w:val"), "0")


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(12.4)

    styles = doc.styles
    title_style = styles.add_style("DanhMucTitle", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = normal
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(13)
    title_style.font.bold = True
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(14)
    title_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    entry_style = styles.add_style("DanhMucEntry", WD_STYLE_TYPE.PARAGRAPH)
    entry_style.base_style = normal
    entry_style.font.name = "Times New Roman"
    entry_style.font.size = Pt(9.5)
    entry_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    entry_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    entry_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    entry_style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    entry_style.paragraph_format.left_indent = Cm(0.45)
    entry_style.paragraph_format.right_indent = Cm(0.45)
    entry_style.paragraph_format.space_before = Pt(0)
    entry_style.paragraph_format.space_after = Pt(0)
    entry_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    entry_style.paragraph_format.line_spacing = Pt(12.4)


def add_title(doc, text, page_break_before=False):
    paragraph = doc.add_paragraph(style="DanhMucTitle")
    paragraph.paragraph_format.page_break_before = page_break_before
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_font(run, "Times New Roman", 13, bold=True)
    set_keep_lines(paragraph)


def add_entry(doc, label, page):
    paragraph = doc.add_paragraph(style="DanhMucEntry")
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(16.55), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    paragraph.add_run(label)
    paragraph.add_run("\t")
    paragraph.add_run(str(page))
    for run in paragraph.runs:
        set_font(run, "Times New Roman", 9.5)
    set_keep_lines(paragraph)


def audit_docx(path):
    document = Document(path)
    titles = [p for p in document.paragraphs if p.style.name == "DanhMucTitle"]
    entries = [p for p in document.paragraphs if p.style.name == "DanhMucEntry"]
    assert len(titles) == 2, f"Expected 2 titles, found {len(titles)}"
    assert len(entries) == len(FIGURES) + len(TABLES), (
        f"Expected {len(FIGURES) + len(TABLES)} entries, found {len(entries)}"
    )
    assert titles[1].paragraph_format.page_break_before is True
    assert all("\t" in p.text for p in entries)
    assert all(p.text.rsplit("\t", 1)[-1].isdigit() for p in entries)

    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
    assert len(re.findall(r'w:leader="dot"', document_xml)) == len(entries)
    assert "Times New Roman" in styles_xml
    assert "DANH MỤC HÌNH VẼ" in document_xml
    assert "DANH MỤC BẢNG BIỂU" in document_xml


def build():
    doc = Document()
    configure_document(doc)

    add_title(doc, "DANH MỤC HÌNH VẼ")
    for label, page in FIGURES:
        add_entry(doc, label, page)

    add_title(doc, "DANH MỤC BẢNG BIỂU", page_break_before=True)
    for label, page in TABLES:
        add_entry(doc, label, page)

    doc.core_properties.title = "Danh mục hình vẽ và danh mục bảng biểu"
    doc.core_properties.subject = "Tổng hợp từ báo cáo Custom Keyboard Builder"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    audit_docx(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Figures: {len(FIGURES)}")
    print(f"Tables: {len(TABLES)}")


if __name__ == "__main__":
    build()
