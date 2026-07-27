from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document: Document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield "table", Table(child, document)


def rgb_string(run) -> str | None:
    color = run.font.color.rgb
    return str(color) if color is not None else None


def inspect(input_path: Path, output_path: Path) -> None:
    doc = Document(input_path)
    blocks = []
    red_notes = []
    headings = []
    figure_paragraphs = []

    paragraph_index = 0
    table_index = 0
    for block_index, (kind, block) in enumerate(iter_blocks(doc)):
        if kind == "paragraph":
            paragraph_index += 1
            p = block
            run_info = []
            red_text_parts = []
            for run in p.runs:
                color = rgb_string(run)
                info = {
                    "text": run.text,
                    "color": color,
                    "highlight": str(run.font.highlight_color) if run.font.highlight_color is not None else None,
                    "bold": run.bold,
                    "italic": run.italic,
                    "size_pt": run.font.size.pt if run.font.size else None,
                    "font": run.font.name,
                }
                run_info.append(info)
                if color and color.upper() in {"FF0000", "C00000", "E60000", "D00000"}:
                    red_text_parts.append(run.text)

            has_drawing = bool(p._p.xpath(".//w:drawing | .//w:pict"))
            has_page_break = bool(p._p.xpath(".//w:br[@w:type='page']"))
            style_name = p.style.name if p.style else ""
            record = {
                "block_index": block_index,
                "kind": kind,
                "paragraph_index": paragraph_index,
                "style": style_name,
                "text": p.text,
                "alignment": str(p.alignment) if p.alignment is not None else None,
                "page_break_before": p.paragraph_format.page_break_before,
                "keep_with_next": p.paragraph_format.keep_with_next,
                "has_page_break": has_page_break,
                "has_drawing": has_drawing,
                "bookmarks": [
                    node.get(qn("w:name"))
                    for node in p._p.xpath(".//w:bookmarkStart")
                    if node.get(qn("w:name"))
                ],
                "runs": run_info,
            }
            blocks.append(record)
            if style_name.lower().startswith("heading") or style_name.lower().startswith("tiêu đề"):
                headings.append({"paragraph_index": paragraph_index, "style": style_name, "text": p.text})
            if red_text_parts:
                red_notes.append(
                    {
                        "paragraph_index": paragraph_index,
                        "full_text": p.text,
                        "red_text": "".join(red_text_parts),
                    }
                )
            if has_drawing:
                figure_paragraphs.append({"paragraph_index": paragraph_index, "text": p.text})
        else:
            table_index += 1
            table = block
            rows = []
            for row in table.rows:
                rows.append([cell.text for cell in row.cells])
            blocks.append(
                {
                    "block_index": block_index,
                    "kind": kind,
                    "table_index": table_index,
                    "style": table.style.name if table.style else "",
                    "rows": rows,
                }
            )

    zip_info = {}
    with zipfile.ZipFile(input_path) as zf:
        names = set(zf.namelist())
        document_xml = zf.read("word/document.xml")
        zip_info = {
            "has_comments": "word/comments.xml" in names,
            "has_footnotes": "word/footnotes.xml" in names,
            "tracked_insertions": len(re.findall(br"<w:ins(?:\s|>)", document_xml)),
            "tracked_deletions": len(re.findall(br"<w:del(?:\s|>)", document_xml)),
            "field_count": document_xml.count(b"<w:fldChar"),
            "bookmark_count": document_xml.count(b"<w:bookmarkStart"),
            "image_parts": len([n for n in names if n.startswith("word/media/")]),
        }

    result = {
        "input": str(input_path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "sections": [
            {
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
            }
            for section in doc.sections
        ],
        "zip_info": zip_info,
        "headings": headings,
        "red_notes": red_notes,
        "figure_paragraphs": figure_paragraphs,
        "blocks": blocks,
    }
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    outline_lines = []
    for block in blocks:
        if block["kind"] == "paragraph":
            flags = []
            if block["has_drawing"]:
                flags.append("IMAGE")
            if block["has_page_break"] or block["page_break_before"]:
                flags.append("PAGE_BREAK")
            flags.extend(f"BOOKMARK={name}" for name in block.get("bookmarks", []))
            flag_text = f" [{' '.join(flags)}]" if flags else ""
            outline_lines.append(
                f"P{block['paragraph_index']:03d} | {block['style']} | {block['text']}{flag_text}"
            )
        else:
            outline_lines.append(f"T{block['table_index']:03d} | {block['style']}")
            for row in block["rows"]:
                outline_lines.append("    | " + " || ".join(cell.replace("\n", " ") for cell in row))
    output_path.with_suffix(".txt").write_text("\n".join(outline_lines), encoding="utf-8")


if __name__ == "__main__":
    inspect(Path(sys.argv[1]), Path(sys.argv[2]))
