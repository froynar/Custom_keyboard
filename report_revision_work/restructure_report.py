from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


FONT_NAME = "Times New Roman"


TOC_ENTRIES = [
    ("opening", 1, "MỞ ĐẦU"),
    ("opening_1", 2, "1. Giới thiệu và lý do chọn đề tài"),
    ("opening_2", 2, "2. Mục đích và nhiệm vụ của đề tài"),
    ("opening_3", 2, "3. Phân công nhiệm vụ của nhóm"),
    ("opening_4", 2, "4. Nội dung sơ lược của báo cáo"),
    ("ch1", 1, "CHƯƠNG 1. TỔNG QUAN HỆ THỐNG CUSTOM KEYBOARD BUILDER"),
    ("ch1_1", 2, "1.1. Bối cảnh và vấn đề cần giải quyết"),
    ("ch1_data", 2, "1.2. Yêu cầu về dữ liệu"),
    ("ch1_2", 2, "1.3. Mục tiêu và phạm vi hệ thống"),
    ("ch1_3", 2, "1.4. Tác nhân và yêu cầu chức năng"),
    ("ch1_4", 2, "1.5. Yêu cầu phi chức năng"),
    ("ch1_tech", 2, "1.6. Công nghệ và môi trường phát triển"),
    ("ch1_tech_1", 3, "1.6.1. Ngôn ngữ C# và nền tảng .NET 10"),
    ("ch1_tech_2", 3, "1.6.2. WPF và XAML"),
    ("ch1_tech_3", 3, "1.6.3. Microsoft SQL Server Express và SSMS"),
    ("ch1_tech_4", 3, "1.6.4. Microsoft Visual Studio"),
    ("ch1_tech_5", 3, "1.6.5. Keyboard kit và phạm vi phần cứng"),
    ("ch1_lib", 2, "1.7. Thư viện và thành phần kỹ thuật"),
    ("ch1_lib_1", 3, "1.7.1. Microsoft.Data.SqlClient 6.1.1"),
    ("ch1_lib_2", 3, "1.7.2. LiveChartsCore.SkiaSharpView.WPF 2.0.4"),
    ("ch1_lib_3", 3, "1.7.3. MQTTnet 4.3.7.1207"),
    ("ch1_lib_4", 3, "1.7.4. Bảo mật mật khẩu và đa ngôn ngữ"),
    ("ch1_arch", 2, "1.8. Kiến trúc phần mềm"),
    ("ch1_arch_1", 3, "1.8.1. Mô hình Model-View-ViewModel"),
    ("ch1_arch_2", 3, "1.8.2. Lớp Service"),
    ("ch1_arch_3", 3, "1.8.3. Lớp Repository"),
    ("ch1_arch_4", 3, "1.8.4. Hạ tầng Data/Realtime và DeviceSimulator"),
    ("ch1_arch_5", 3, "1.8.5. Composition root và điều hướng theo vai trò"),
    ("ch1_9", 2, "1.9. Kết luận chương"),
    ("ch2", 1, "CHƯƠNG 2. PHÂN TÍCH THIẾT KẾ HỆ THỐNG"),
    ("ch2_1", 2, "2.1. Use Case Diagram"),
    ("ch2_2", 2, "2.2. Activity Diagram"),
    ("ch2_3", 2, "2.3. Sequence Diagram"),
    ("ch2_4", 2, "2.4. Entity-Relationship Diagram (ERD)"),
    ("ch2_4_1", 3, "2.4.1. Sơ đồ ERD tổng quát"),
    ("ch2_4_2", 3, "2.4.2. Sơ đồ ERD chi tiết"),
    ("ch2_5", 2, "2.5. Đối chiếu diagram với thành phần triển khai"),
    ("ch2_6", 2, "2.6. Kết luận chương"),
    ("ch3", 1, "CHƯƠNG 3. TRIỂN KHAI VÀ KIỂM THỬ"),
    ("ch3_1", 2, "3.1. Triển khai cơ sở dữ liệu SQL Server"),
    ("ch3_1_1", 3, "3.1.1. Nhóm tài khoản và phân quyền"),
    ("ch3_1_2", 3, "3.1.2. Nhóm danh mục linh kiện"),
    ("ch3_1_3", 3, "3.1.3. Nhóm build và request"),
    ("ch3_1_4", 3, "3.1.4. Nhóm thiết bị và kiểm tra QC"),
    ("ch3_1_5", 3, "3.1.5. Nhóm audit và chat"),
    ("ch3_2", 2, "3.2. Triển khai giao diện tiêu biểu"),
    ("ch3_3", 2, "3.3. Kiểm thử và xác minh"),
    ("ch3_4", 2, "3.4. Đánh giá mức độ hoàn thành"),
    ("ch3_5", 2, "3.5. Kết luận chương"),
    ("conclusion", 1, "KẾT LUẬN"),
    ("conclusion_1", 2, "1. Mức độ hoàn thành nhiệm vụ"),
    ("conclusion_2", 2, "2. Mức độ phù hợp với mục đích đề tài"),
    ("conclusion_3", 2, "3. Hạn chế còn tồn tại"),
    ("conclusion_4", 2, "4. Hướng phát triển"),
]


FIGURE_ENTRIES = [
    ("fig_2_1", "Hình 2.1. Sơ đồ Use Case tổng hợp của hệ thống Custom Keyboard Builder"),
    ("fig_2_2", "Hình 2.2. Sơ đồ Activity tổng hợp của các luồng nghiệp vụ chính"),
    ("fig_2_3", "Hình 2.3. Sơ đồ Sequence tổng hợp của các luồng tương tác chính"),
    ("fig_2_4", "Hình 2.4. Sơ đồ ERD tổng quát của hệ thống Custom Keyboard Builder"),
    ("fig_2_5", "Hình 2.5. Sơ đồ ERD chi tiết gồm 21 bảng và 33 khóa ngoại"),
    ("fig_3_1", "Hình 3.1. Giao diện đăng nhập của ứng dụng"),
    ("fig_3_register", "Hình 3.2. Giao diện đăng ký tài khoản Buyer"),
    ("fig_3_2", "Hình 3.3. Giao diện Buyer Dashboard và khu vực cấu hình build"),
    ("fig_3_seller", "Hình 3.4. Giao diện Seller Dashboard và xử lý yêu cầu lắp ráp"),
    ("fig_3_admin", "Hình 3.5. Giao diện Admin Dashboard và quản trị hệ thống"),
    ("fig_3_chat", "Hình 3.6. Giao diện trao đổi tin nhắn theo vai trò"),
    ("fig_3_user_menu", "Hình 3.7. Giao diện hồ sơ và menu người dùng"),
]


TABLE_ENTRIES = [
    ("tbl_0_1", "Bảng 0.1. Phân công nhiệm vụ nhóm"),
    ("tbl_1_data", "Bảng 1.1. Yêu cầu về dữ liệu"),
    ("tbl_1_1", "Bảng 1.2. Tác nhân và mục tiêu sử dụng hệ thống"),
    ("tbl_1_2", "Bảng 1.3. Yêu cầu chức năng chính"),
    ("tbl_1_3", "Bảng 1.4. Yêu cầu phi chức năng"),
    ("tbl_2_1", "Bảng 2.1. Nhóm bảng trong cơ sở dữ liệu"),
    ("tbl_2_2", "Bảng 2.2. Đối chiếu bốn loại diagram với artifact nguồn"),
    ("tbl_3_1", "Bảng 3.1. Thứ tự triển khai 21 bảng theo nhóm phụ thuộc"),
    ("tbl_3_2", "Bảng 3.2. Đối chiếu giao diện với mã nguồn"),
    ("tbl_3_3", "Bảng 3.3. Kết quả build và xác minh tại thời điểm lập báo cáo"),
    ("tbl_3_4", "Bảng 3.4. Đánh giá mức độ hoàn thành theo mục tiêu"),
]


def set_run_font(run, size: float = 13, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = FONT_NAME
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT_NAME)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.highlight_color = None


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text.strip():
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_paragraph_startswith(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found with prefix: {prefix}")


def find_heading_startswith(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading") and paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Heading not found with prefix: {prefix}")


def clear_paragraph_content(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag in {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
            continue
        paragraph._p.remove(child)


def replace_paragraph_text(
    paragraph: Paragraph,
    text: str,
    *,
    size: float = 13,
    bold: bool | None = None,
    alignment=None,
) -> Paragraph:
    clear_paragraph_content(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if alignment is not None:
        paragraph.alignment = alignment
    return paragraph


def max_bookmark_id(doc: Document) -> int:
    values = []
    for node in doc.element.xpath(".//w:bookmarkStart"):
        value = node.get(qn("w:id"))
        if value and value.isdigit():
            values.append(int(value))
    return max(values, default=0)


class BookmarkManager:
    def __init__(self, doc: Document) -> None:
        self.doc = doc
        self.next_id = max_bookmark_id(doc) + 1

    def names(self) -> set[str]:
        return {
            node.get(qn("w:name"))
            for node in self.doc.element.xpath(".//w:bookmarkStart")
            if node.get(qn("w:name"))
        }

    def add(self, paragraph: Paragraph, name: str) -> None:
        if name in self.names():
            return
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(self.next_id))
        bookmark_start.set(qn("w:name"), name)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), str(self.next_id))
        insert_at = 1 if paragraph._p.pPr is not None else 0
        paragraph._p.insert(insert_at, bookmark_start)
        paragraph._p.append(bookmark_end)
        self.next_id += 1


def rename_bookmark(paragraph: Paragraph, old_name: str, new_name: str) -> None:
    for node in paragraph._p.xpath(".//w:bookmarkStart"):
        if node.get(qn("w:name")) == old_name:
            node.set(qn("w:name"), new_name)


def move_paragraph_before(paragraph: Paragraph, target: Paragraph) -> Paragraph:
    target._p.addprevious(paragraph._p)
    return paragraph


def add_heading_before(
    doc: Document,
    manager: BookmarkManager,
    target: Paragraph,
    text: str,
    level: int,
    anchor: str,
) -> Paragraph:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = False
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 14 if level == 2 else 13, bold=True)
    manager.add(paragraph, anchor)
    return move_paragraph_before(paragraph, target)


def add_body_before(doc: Document, target: Paragraph, text: str) -> Paragraph:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    set_run_font(run, size=13)
    return move_paragraph_before(paragraph, target)


def add_bullet_before(doc: Document, target: Paragraph, text: str) -> Paragraph:
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, size=13)
    return move_paragraph_before(paragraph, target)


def add_internal_hyperlink(paragraph: Paragraph, text: str, anchor: str, *, size: float, bold: bool) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)
    r_pr.append(fonts)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_node)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), "000000")
    r_pr.append(color_node)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    r_pr.append(underline)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def replace_paragraph_with_reference(
    paragraph: Paragraph,
    prefix: str,
    link_text: str,
    anchor: str,
    suffix: str = ".",
) -> None:
    clear_paragraph_content(paragraph)
    prefix_run = paragraph.add_run(prefix)
    set_run_font(prefix_run, size=12)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)
    r_pr.append(fonts)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), "24")
    r_pr.append(size_node)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), "0563C1")
    r_pr.append(color_node)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = link_text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

    suffix_run = paragraph.add_run(suffix)
    set_run_font(suffix_run, size=12)


def add_nav_line_before(
    doc: Document,
    target: Paragraph,
    text: str,
    anchor: str,
    page: str,
    *,
    level: int = 1,
    size: float = 12,
    bold: bool = False,
) -> Paragraph:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.25 * (level - 1))
    paragraph.paragraph_format.first_line_indent = Inches(-0.08 if level > 1 else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.35), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    add_internal_hyperlink(paragraph, text, anchor, size=size, bold=bold)
    page_run = paragraph.add_run(f"\t{page}")
    set_run_font(page_run, size=size, bold=bold)
    return move_paragraph_before(paragraph, target)


def remove_between(start: Paragraph, end: Paragraph) -> None:
    current = start._p.getnext()
    while current is not None and current is not end._p:
        following = current.getnext()
        current.getparent().remove(current)
        current = following


def remove_range_inclusive(start: Paragraph, end_exclusive: Paragraph) -> None:
    current = start._p
    while current is not None and current is not end_exclusive._p:
        following = current.getnext()
        current.getparent().remove(current)
        current = following


def remove_highlights(doc: Document) -> None:
    for root in [doc.element, *[section.header._element for section in doc.sections], *[section.footer._element for section in doc.sections]]:
        for node in list(root.xpath(".//w:highlight")):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def clear_assignment_table(doc: Document) -> None:
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "Thành viên":
            for row in table.rows[1:]:
                for cell in row.cells:
                    cell.text = ""
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
            return
    raise ValueError("Assignment table not found")


def ensure_target_bookmarks(doc: Document, manager: BookmarkManager) -> None:
    exact_targets = {
        "MỞ ĐẦU": "opening",
        "1. Giới thiệu và lý do chọn đề tài": "opening_1",
        "2. Mục đích và nhiệm vụ của đề tài": "opening_2",
        "3. Phân công nhiệm vụ của nhóm": "opening_3",
        "4. Nội dung sơ lược của báo cáo": "opening_4",
        "CHƯƠNG 1. TỔNG QUAN HỆ THỐNG CUSTOM KEYBOARD BUILDER": "ch1",
        "1.1. Bối cảnh và vấn đề cần giải quyết": "ch1_1",
        "1.2. Yêu cầu về dữ liệu": "ch1_data",
        "1.3. Mục tiêu và phạm vi hệ thống": "ch1_2",
        "1.4. Tác nhân và yêu cầu chức năng": "ch1_3",
        "1.5. Yêu cầu phi chức năng": "ch1_4",
        "1.9. Kết luận chương": "ch1_9",
        "CHƯƠNG 2. PHÂN TÍCH THIẾT KẾ HỆ THỐNG": "ch2",
        "2.1. Use Case Diagram": "ch2_1",
        "2.2. Activity Diagram": "ch2_2",
        "2.3. Sequence Diagram": "ch2_3",
        "2.4. Entity-Relationship Diagram (ERD)": "ch2_4",
        "2.4.1. Sơ đồ ERD tổng quát": "ch2_4_1",
        "2.4.2. Sơ đồ ERD chi tiết": "ch2_4_2",
        "2.5. Đối chiếu diagram với thành phần triển khai": "ch2_5",
        "2.6. Kết luận chương": "ch2_6",
    }
    for text, anchor in exact_targets.items():
        manager.add(find_paragraph(doc, text), anchor)


def remove_named_bookmark(doc: Document, name: str) -> None:
    bookmark_ids = []
    for node in list(doc.element.xpath(".//w:bookmarkStart")):
        if node.get(qn("w:name")) == name:
            bookmark_ids.append(node.get(qn("w:id")))
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    for node in list(doc.element.xpath(".//w:bookmarkEnd")):
        if node.get(qn("w:id")) in bookmark_ids:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def reanchor_navigation_targets(doc: Document, manager: BookmarkManager) -> None:
    entries = [(anchor, text) for anchor, _level, text in TOC_ENTRIES]
    entries.extend(FIGURE_ENTRIES)
    entries.extend(TABLE_ENTRIES)
    for anchor, text in entries:
        target = find_paragraph(doc, text)
        remove_named_bookmark(doc, anchor)
        manager.add(target, anchor)


def rebuild_front_matter(doc: Document) -> None:
    toc_title = find_paragraph(doc, "MỤC LỤC")
    figure_title = find_paragraph(doc, "DANH MỤC HÌNH VẼ")
    table_title = find_paragraph(doc, "DANH MỤC BẢNG BIỂU")
    abbreviation_title = find_paragraph(doc, "DANH MỤC KÝ HIỆU VÀ CHỮ VIẾT TẮT")

    remove_between(toc_title, figure_title)
    for anchor, level, text in TOC_ENTRIES:
        add_nav_line_before(doc, figure_title, text, anchor, "00", level=level, bold=level == 1)

    remove_between(figure_title, table_title)
    for anchor, text in FIGURE_ENTRIES:
        add_nav_line_before(doc, table_title, text, anchor, "00", level=1, size=11.5)

    remove_between(table_title, abbreviation_title)
    for anchor, text in TABLE_ENTRIES:
        add_nav_line_before(doc, abbreviation_title, text, anchor, "00", level=1, size=11.5)


def add_missing_abbreviations(doc: Document) -> None:
    opening = find_paragraph(doc, "MỞ ĐẦU")
    existing = {paragraph.text.strip() for paragraph in doc.paragraphs}
    entries = [
        "ADO.NET: Tập API truy cập dữ liệu của nền tảng .NET.",
        "IDE: Integrated Development Environment - môi trường phát triển tích hợp.",
        "PBKDF2: Password-Based Key Derivation Function 2 - hàm dẫn xuất khóa dùng để băm mật khẩu.",
    ]
    for text in entries:
        if text not in existing:
            add_bullet_before(doc, opening, text)


def main(input_path: Path, output_path: Path) -> None:
    doc = Document(input_path)
    remove_highlights(doc)
    manager = BookmarkManager(doc)

    # Mở đầu: hoàn thiện mục đích/nhiệm vụ và thực hiện note để trống bảng phân công.
    replace_paragraph_text(
        find_paragraph(doc, "1. Giới thiệu và lý do chọn đề tài"),
        "1. Giới thiệu và lý do chọn đề tài",
        size=14,
        bold=True,
    )
    opening_paragraphs = [
        (
            "Hiện nay, nhu cầu sử dụng bàn phím cơ tự lắp ráp và cá nhân hóa ngày càng tăng, đặc biệt đối với người làm việc thường xuyên với máy tính như lập trình viên, nhân viên văn phòng, người chơi trò chơi điện tử và người dùng phổ thông. Một bộ bàn phím tùy chỉnh cho phép lựa chọn đặc tính gõ, độ ồn, bố cục và hình thức phù hợp với nhu cầu riêng.",
            "Hiện nay, việc sử dụng bàn phím cơ tự xây dựng",
        ),
        (
            "Tuy nhiên, số lượng loại linh kiện và tiêu chuẩn tương thích khá lớn. Việc lựa chọn riêng lẻ keyboard kit, switch, keycap, stabilizer và phụ kiện có thể gây nhầm lẫn, làm tăng thời gian tìm hiểu và phát sinh cấu hình không phù hợp. Nhu cầu quản lý tập trung cấu hình, giá tại thời điểm đặt hàng và quá trình xử lý vì vậy được đặt ra.",
            "Tuy nhiên, việc tự tìm hiểu",
        ),
        (
            "Từ bối cảnh trên, hệ thống Custom Keyboard Builder được xây dựng để hỗ trợ tạo cấu hình bàn phím theo mô hình kit-based, gửi yêu cầu lắp ráp cho Seller, theo dõi trạng thái, lưu kết quả kiểm tra chất lượng và quản trị dữ liệu tập trung. Ứng dụng desktop WPF phục vụ ba vai trò Buyer, Seller và Admin; SQL Server được sử dụng làm nguồn dữ liệu chính.",
            "Dự án Custom Keyboard Builder được xây dựng",
        ),
    ]
    for replacement, prefix in opening_paragraphs:
        replace_paragraph_text(find_paragraph_startswith(doc, prefix), replacement)

    purpose_heading = find_paragraph(doc, "2. Mục đích của đề tài")
    replace_paragraph_text(
        purpose_heading,
        "2. Mục đích và nhiệm vụ của đề tài",
        size=14,
        bold=True,
    )
    purpose_body = find_paragraph_startswith(doc, "Mục đích của đề tài là xây dựng")
    replace_paragraph_text(
        purpose_body,
        "Mục đích của đề tài là xây dựng một ứng dụng quản lý cấu hình và yêu cầu lắp ráp bàn phím tùy chỉnh. Hệ thống hướng tới việc giảm khó khăn khi lựa chọn linh kiện, bảo đảm dữ liệu cấu hình được lưu nhất quán, hỗ trợ phối hợp giữa Buyer và Seller, đồng thời cung cấp công cụ quản trị và truy vết cho Admin.",
    )
    assignment_heading = find_heading_startswith(doc, "3. Phân công nhiệm vụ của nhóm")
    add_body_before(doc, assignment_heading, "Để đạt được mục đích trên, các nhiệm vụ chính được xác định như sau:")
    for item in [
        "Khảo sát bài toán, xác định tác nhân, yêu cầu chức năng, yêu cầu phi chức năng và phạm vi thực hiện.",
        "Phân tích và thiết kế hệ thống thông qua năm sơ đồ: Use Case, Activity, Sequence, ERD tổng quát và ERD chi tiết.",
        "Thiết kế cơ sở dữ liệu SQL Server, các ràng buộc toàn vẹn và dữ liệu phục vụ kiểm thử.",
        "Triển khai WPF theo MVVM, kết hợp Service, Repository, MQTT và mô phỏng QC.",
        "Thực hiện build, xác minh các luồng chính, tổng hợp kết quả và hoàn thiện báo cáo.",
    ]:
        add_bullet_before(doc, assignment_heading, item)
    replace_paragraph_text(assignment_heading, "3. Phân công nhiệm vụ của nhóm", size=14, bold=True)
    replace_paragraph_with_reference(
        find_paragraph_startswith(doc, "Các vị trí họ tên và mã số sinh viên"),
        "",
        "Bảng 0.1",
        "tbl_0_1",
        " được bố trí để bổ sung thông tin thành viên và nhiệm vụ được phân công.",
    )
    clear_assignment_table(doc)

    summary = find_paragraph_startswith(doc, "Báo cáo gồm")
    replace_paragraph_text(
        summary,
        "Báo cáo gồm 00 trang, 3 chương chính, 12 hình vẽ, 11 bảng biểu và 25 khối mã nguồn:",
    )
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương 1 Tổng quan hệ thống:"),
        "Chương 1 Tổng quan hệ thống: trình bày bối cảnh, mục tiêu, phạm vi, yêu cầu, công nghệ, thư viện và kiến trúc phần mềm.",
    )
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương 2 Phân tích thiết kế hệ thống:"),
        "Chương 2 Phân tích thiết kế hệ thống: tập trung phân tích năm sơ đồ Use Case, Activity, Sequence, ERD tổng quát và ERD chi tiết.",
    )
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương 3 Triển khai và kiểm thử:"),
        "Chương 3 Triển khai và kiểm thử: trình bày cơ sở dữ liệu, giao diện và mã nguồn tiêu biểu, kết quả kiểm thử và đánh giá mức độ hoàn thành.",
    )

    # Chương 1: chuyển và mở rộng phần công nghệ, thư viện, kiến trúc từ Chương 2.
    chapter_1_intro = find_paragraph_startswith(doc, "Chương này trình bày bối cảnh hình thành")
    replace_paragraph_text(
        chapter_1_intro,
        "Chương này trình bày bối cảnh, mục tiêu, phạm vi, yêu cầu dữ liệu và yêu cầu nghiệp vụ của hệ thống. Các công nghệ, phần mềm, thư viện và kiến trúc được giới thiệu để làm cơ sở cho hoạt động phân tích thiết kế ở Chương 2.",
    )
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Một bộ bàn phím custom yêu cầu nhiều linh kiện"),
        "Một bộ bàn phím tùy chỉnh được hình thành từ nhiều nhóm linh kiện có quan hệ tương thích về công nghệ switch, kiểu mount, layout và số lượng phím. Nếu thông tin được tìm kiếm và kiểm tra thủ công, người dùng mới có thể mất nhiều thời gian hoặc tạo cấu hình không phù hợp. Hệ thống vì vậy được xây dựng để hỗ trợ lựa chọn linh kiện có kiểm soát và lưu lại cấu hình một cách nhất quán.",
    )
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Mô hình hiện tại lấy keyboard kit làm nền tảng"),
        "Mô hình hiện tại lấy keyboard kit làm nền tảng. Một kit bao gồm case, PCB, plate và các phần đi kèm; Buyer tiếp tục chọn switch, keycap, stabilizer, accessory và ghi chú mod. Bảng build_items quy định mỗi dòng chỉ tham chiếu một loại sản phẩm, còn giá snapshot được lưu để giữ nguyên giá trị cấu hình tại thời điểm request được gửi.",
    )
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Dữ liệu của hệ thống cần được lưu tập trung"),
        "Dữ liệu của hệ thống được lưu tập trung trên Microsoft SQL Server Express. SQL Server Management Studio (SSMS) được sử dụng để tạo schema, thực thi script, kiểm tra dữ liệu và theo dõi các ràng buộc trong quá trình phát triển.",
    )

    functional_heading = find_paragraph(doc, "1.4. Tác nhân và yêu cầu chức năng")
    add_body_before(
        doc,
        functional_heading,
        "Phạm vi thực hiện bao gồm quản lý tài khoản và vai trò, danh mục linh kiện, cấu hình build, gửi và xử lý request, đơn đăng ký Seller, chat lưu trong cơ sở dữ liệu, thống kê dashboard, audit log và kiểm tra QC bằng dữ liệu mô phỏng.",
    )
    add_body_before(
        doc,
        functional_heading,
        "Các nội dung ngoài phạm vi hiện tại gồm thanh toán trực tuyến, quản lý vận chuyển, kết nối ESP32 hoặc trạm đo phần cứng thật và chat realtime bằng SignalR. Các nội dung này được xem là hướng mở rộng thay vì điều kiện nghiệm thu của phiên bản hiện tại.",
    )

    conclusion_heading = find_paragraph(doc, "1.6. Kết luận chương")
    add_heading_before(doc, manager, conclusion_heading, "1.6. Công nghệ và môi trường phát triển", 2, "ch1_tech")
    add_body_before(
        doc,
        conclusion_heading,
        "Hệ thống được phát triển dưới dạng ứng dụng desktop trên Windows. Việc lựa chọn công nghệ ưu tiên khả năng tích hợp với WPF, SQL Server và hệ sinh thái .NET, đồng thời phù hợp với mô hình phân lớp được áp dụng trong mã nguồn.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.6.1. Ngôn ngữ C# và nền tảng .NET 10", 3, "ch1_tech_1")
    add_body_before(
        doc,
        conclusion_heading,
        "C# là ngôn ngữ lập trình hướng đối tượng có kiểu dữ liệu chặt chẽ do Microsoft phát triển. Các đặc điểm như đóng gói, kế thừa, đa hình, generic, xử lý ngoại lệ và lập trình bất đồng bộ bằng async/await phù hợp với việc tổ chức nghiệp vụ, truy cập dữ liệu và duy trì khả năng phản hồi của giao diện.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "Dự án sử dụng target framework net10.0-windows của .NET 10. Nền tảng .NET cung cấp runtime, thư viện cơ sở, cơ chế quản lý bộ nhớ và hệ thống package NuGet; hậu tố windows cho phép sử dụng WPF và các API giao diện đặc thù của Windows.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.6.2. WPF và XAML", 3, "ch1_tech_2")
    add_body_before(
        doc,
        conclusion_heading,
        "Windows Presentation Foundation (WPF) là framework xây dựng giao diện desktop trên Windows. WPF cung cấp hệ thống control, DataGrid, style, resource, data binding và command, qua đó hỗ trợ tách phần hiển thị khỏi logic xử lý theo MVVM.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "Extensible Application Markup Language (XAML) được dùng để khai báo bố cục của Login, Buyer Dashboard, Seller Dashboard, Admin Dashboard, Chat và User Menu. Code C# phía ViewModel cung cấp dữ liệu và command; binding giúp giao diện cập nhật theo trạng thái mà không cần truy cập trực tiếp cơ sở dữ liệu.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.6.3. Microsoft SQL Server Express và SSMS", 3, "ch1_tech_3")
    add_body_before(
        doc,
        conclusion_heading,
        "Microsoft SQL Server là hệ quản trị cơ sở dữ liệu quan hệ, chịu trách nhiệm lưu tài khoản, danh mục linh kiện, build, request, chat, audit và kết quả QC. Cấu hình hiện tại trỏ tới instance SQL Server Express KHOADZS1VN\\SQLEXPRESS và cơ sở dữ liệu CustomKeyboard_Refactor; số phiên bản máy chủ cụ thể không bị khóa cứng trong mã nguồn.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "SQL Server Management Studio là công cụ quản trị được sử dụng để chạy DDL, seed data, truy vấn kiểm tra và quan sát quan hệ giữa các bảng. SQL Server là nơi lưu dữ liệu chính; SSMS chỉ là công cụ làm việc với máy chủ, không phải nơi lưu dữ liệu của ứng dụng.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.6.4. Microsoft Visual Studio", 3, "ch1_tech_4")
    add_body_before(
        doc,
        conclusion_heading,
        "Microsoft Visual Studio là môi trường phát triển tích hợp (IDE) phục vụ soạn thảo C# và XAML, quản lý project SDK-style, khôi phục package NuGet, build, gỡ lỗi và theo dõi lỗi biên dịch. Dự án cũng có thể được build bằng lệnh dotnet khi máy đã cài .NET SDK tương ứng.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.6.5. Keyboard kit và phạm vi phần cứng", 3, "ch1_tech_5")
    add_body_before(
        doc,
        conclusion_heading,
        "Khái niệm keyboard kit trong đề tài là bộ sản phẩm nền gồm case, PCB, plate và các phần đi kèm; đây không phải development kit hoặc bo mạch điều khiển dùng để phát triển phần cứng. Buyer lựa chọn kit từ catalog rồi bổ sung switch, keycap, stabilizer, accessory và mod để tạo build.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "Phiên bản hiện tại không sử dụng ESP32, Arduino hoặc thiết bị nhúng thật. Quá trình kiểm tra từng phím được mô phỏng bằng DeviceSimulator; hướng kết nối trạm QC vật lý chỉ được xem là khả năng mở rộng sau phạm vi đề tài.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.7. Thư viện và thành phần kỹ thuật", 2, "ch1_lib")
    add_body_before(
        doc,
        conclusion_heading,
        "Các package NuGet và thành phần kỹ thuật bổ sung được lựa chọn theo nhu cầu truy cập dữ liệu, trực quan hóa dashboard, truyền thông realtime, bảo mật và đa ngôn ngữ.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.7.1. Microsoft.Data.SqlClient 6.1.1", 3, "ch1_lib_1")
    add_body_before(
        doc,
        conclusion_heading,
        "Microsoft.Data.SqlClient là data provider thuộc ADO.NET dành cho SQL Server. Thư viện cung cấp SqlConnection, SqlCommand, SqlParameter, SqlDataReader và transaction. Các Repository sử dụng câu lệnh có tham số để giảm nguy cơ SQL injection và ánh xạ dữ liệu trả về thành Model của hệ thống.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.7.2. LiveChartsCore.SkiaSharpView.WPF 2.0.4", 3, "ch1_lib_2")
    add_body_before(
        doc,
        conclusion_heading,
        "LiveChartsCore kết hợp SkiaSharp và WPF để hiển thị biểu đồ KPI, doanh thu, trạng thái request và dữ liệu phân tích trên dashboard. Dữ liệu biểu đồ được ViewModel nhận từ StatsService thay vì được truy vấn trực tiếp trong View.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.7.3. MQTTnet 4.3.7.1207", 3, "ch1_lib_3")
    add_body_before(
        doc,
        conclusion_heading,
        "MQTTnet là thư viện .NET triển khai giao thức MQTT theo mô hình publish/subscribe. Trong hệ thống, MQTT được dùng tùy chọn để phát thông báo request, cập nhật trạng thái và truyền telemetry QC qua broker mặc định localhost:1883. Cơ chế được triển khai theo best-effort; khi broker không khả dụng, luồng nghiệp vụ chính vẫn tiếp tục ở chế độ DB-only.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.7.4. Bảo mật mật khẩu và đa ngôn ngữ", 3, "ch1_lib_4")
    add_body_before(
        doc,
        conclusion_heading,
        "Mật khẩu được xử lý bằng PBKDF2-SHA256 của thư viện mật mã tích hợp trong .NET, với 100.000 vòng lặp, salt ngẫu nhiên 16 byte và hash 32 byte. Giá trị mật khẩu rõ không được lưu trong SQL Server. Phần giao diện sử dụng resource và lớp localization nội bộ để chuyển đổi nội dung giữa tiếng Việt và tiếng Anh.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.8. Kiến trúc phần mềm", 2, "ch1_arch")
    add_body_before(
        doc,
        conclusion_heading,
        "Kiến trúc ứng dụng được tổ chức theo MVVM và mở rộng bằng các lớp Service, Repository, Data/Realtime. Cách phân chia này giúp mỗi nhóm lớp có trách nhiệm rõ ràng và hạn chế việc trộn mã giao diện, quy tắc nghiệp vụ và câu lệnh SQL trong cùng một thành phần.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.8.1. Mô hình Model-View-ViewModel", 3, "ch1_arch_1")
    add_body_before(
        doc,
        conclusion_heading,
        "Model nằm trong Models/**/*.cs, biểu diễn dữ liệu và đối tượng nghiệp vụ như User, KeyboardBuild, BuildRequest, linh kiện, hội thoại và kết quả QC. Model được Repository tạo từ dữ liệu SQL Server, được Service xử lý và được ViewModel sử dụng để cung cấp dữ liệu cho giao diện.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "View nằm trong Views/*.xaml, chịu trách nhiệm khai báo bố cục, control, binding, command, DataGrid, biểu đồ và trạng thái hiển thị. View không trực tiếp viết câu lệnh SQL hoặc quyết định quy tắc chuyển trạng thái nghiệp vụ.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "ViewModel nằm trong ViewModels/*ViewModel.cs, giữ state của màn hình, cung cấp command bất đồng bộ và điều phối Service. Khi thuộc tính trong ViewModel thay đổi, cơ chế binding thông báo cho View cập nhật nội dung tương ứng.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.8.2. Lớp Service", 3, "ch1_arch_2")
    add_body_before(
        doc,
        conclusion_heading,
        "Service nằm trong Services/*.cs và thực thi các trường hợp sử dụng của hệ thống. Các lớp này kiểm tra dữ liệu đầu vào, tính tương thích linh kiện, phân quyền, state machine của request, điều kiện hoàn thành QC và các quy tắc nghiệp vụ trước khi yêu cầu Repository đọc hoặc ghi dữ liệu.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.8.3. Lớp Repository", 3, "ch1_arch_3")
    add_body_before(
        doc,
        conclusion_heading,
        "Repository gồm các interface trong Repositories và phần hiện thực SQL Server trong Repositories/SqlServer. Repository mở kết nối, thực hiện câu lệnh SQL có tham số, quản lý transaction khi cần và ánh xạ kết quả thành Model. Nhờ đó, ViewModel và Service không phải phụ thuộc trực tiếp vào chi tiết truy vấn SQL.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.8.4. Hạ tầng Data/Realtime và DeviceSimulator", 3, "ch1_arch_4")
    add_body_before(
        doc,
        conclusion_heading,
        "Data/SqlServer cung cấp cấu hình, connection factory và kiểm tra schema. Realtime triển khai MQTT phục vụ thông báo và telemetry. DeviceSimulator tạo dữ liệu QC; nếu MQTT không khả dụng, kết quả vẫn được xử lý và lưu trong SQL Server.",
    )

    add_heading_before(doc, manager, conclusion_heading, "1.8.5. Composition root và điều hướng theo vai trò", 3, "ch1_arch_5")
    add_body_before(
        doc,
        conclusion_heading,
        "MainWindow là composition root, tức điểm lắp ráp thành phần khi ứng dụng khởi động. Tại đây, connection factory, Repository, Service, MQTT, DeviceSimulator và MainShellViewModel được khởi tạo, cấu hình và truyền phụ thuộc; MainWindow không thực thi nghiệp vụ thay Service.",
    )
    add_body_before(
        doc,
        conclusion_heading,
        "Sau đăng nhập, MainShellViewModel kiểm tra role và điều hướng tới dashboard Buyer, Seller hoặc Admin. Luồng chính đi theo View -> ViewModel -> Service -> Repository -> SQL Server; MQTT chỉ bổ sung thông báo gần thời gian thực.",
    )

    rename_bookmark(conclusion_heading, "ch1_6", "ch1_9")
    replace_paragraph_text(conclusion_heading, "1.9. Kết luận chương", size=14, bold=True)
    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương 1 đã xác định được bài toán"),
        "Chương 1 đã xác định bối cảnh, mục tiêu, phạm vi, yêu cầu, công nghệ, thư viện và kiến trúc của hệ thống. Trên cơ sở đó, Chương 2 mô hình hóa các yêu cầu bằng năm sơ đồ thiết kế.",
    )

    # Chương 2: bỏ phần công nghệ/kiến trúc cũ và chỉ giữ nội dung phân tích thiết kế.
    old_tech_heading = find_paragraph(doc, "2.1. Công nghệ và môi trường triển khai")
    use_case_heading = find_paragraph(doc, "2.3. Use Case Diagram")
    remove_range_inclusive(old_tech_heading, use_case_heading)

    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương này trình bày công nghệ, kiến trúc phần mềm"),
        "Chương này tập trung phân tích và thiết kế hệ thống thông qua năm sơ đồ: Use Case, Activity, Sequence, ERD tổng quát và ERD chi tiết. Các sơ đồ lần lượt mô tả phạm vi chức năng, luồng nghiệp vụ, thứ tự tương tác và cấu trúc dữ liệu của hệ thống.",
    )

    chapter_2_updates = [
        ("2.3. Use Case Diagram", "2.1. Use Case Diagram", "ch2_3", "ch2_1"),
        ("2.4. Activity Diagram", "2.2. Activity Diagram", "ch2_4", "ch2_2"),
        ("2.5. Sequence Diagram", "2.3. Sequence Diagram", "ch2_5", "ch2_3"),
        ("2.6. Entity-Relationship Diagram (ERD)", "2.4. Entity-Relationship Diagram (ERD)", "ch2_6", "ch2_4"),
        ("2.6.1. Sơ đồ ERD tổng quát", "2.4.1. Sơ đồ ERD tổng quát", "ch2_6_1", "ch2_4_1"),
        ("2.6.2. Sơ đồ ERD chi tiết", "2.4.2. Sơ đồ ERD chi tiết", "ch2_6_2", "ch2_4_2"),
        ("2.7. Đối chiếu diagram với thành phần triển khai", "2.5. Đối chiếu diagram với thành phần triển khai", "ch2_7", "ch2_5"),
        ("2.8. Kết luận chương", "2.6. Kết luận chương", "ch2_8", "ch2_6"),
    ]
    for old_text, new_text, old_anchor, new_anchor in chapter_2_updates:
        paragraph = find_paragraph(doc, old_text)
        rename_bookmark(paragraph, old_anchor, new_anchor)
        level = 3 if new_text.startswith("2.4.") and new_text[4:5].isdigit() else 2
        replace_paragraph_text(paragraph, new_text, size=13 if level == 3 else 14, bold=True)
        paragraph.paragraph_format.page_break_before = False
        paragraph.paragraph_format.keep_with_next = True

    database_caption = find_paragraph(doc, "Bảng 2.3. Nhóm bảng trong cơ sở dữ liệu")
    rename_bookmark(database_caption, "tbl_2_3", "tbl_2_1")
    replace_paragraph_text(
        database_caption,
        "Bảng 2.1. Nhóm bảng trong cơ sở dữ liệu",
        size=12,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    database_caption.paragraph_format.keep_with_next = True
    replace_paragraph_with_reference(
        find_paragraph_startswith(doc, "Các nhóm bảng được tổng hợp tại Bảng 2.3"),
        "Các nhóm bảng được tổng hợp tại ",
        "Bảng 2.1",
        "tbl_2_1",
    )

    mapping_caption = find_paragraph(doc, "Bảng 2.4. Đối chiếu bốn loại diagram với artifact nguồn")
    rename_bookmark(mapping_caption, "tbl_2_4", "tbl_2_2")
    replace_paragraph_text(
        mapping_caption,
        "Bảng 2.2. Đối chiếu bốn loại diagram với artifact nguồn",
        size=12,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    mapping_caption.paragraph_format.keep_with_next = True
    replace_paragraph_with_reference(
        find_paragraph_startswith(doc, "Nguồn của năm sơ đồ thuộc bốn nhóm ký pháp được đối chiếu tại Bảng 2.4"),
        "Nguồn của năm sơ đồ thuộc bốn nhóm ký pháp được đối chiếu tại ",
        "Bảng 2.2",
        "tbl_2_2",
    )

    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương 2 đã xác định công nghệ, kiến trúc"),
        "Chương 2 đã mô hình hóa phạm vi chức năng, luồng nghiệp vụ, thứ tự tương tác và cấu trúc dữ liệu thông qua năm sơ đồ thuộc bốn nhóm ký pháp. ERD chi tiết là cơ sở trực tiếp để triển khai 21 bảng SQL Server; các sơ đồ còn lại định hướng việc hiện thực giao diện và nghiệp vụ được trình bày ở Chương 3.",
    )

    replace_paragraph_text(
        find_paragraph_startswith(doc, "Chương 3 đã trình bày 21 DDL"),
        "Chương 3 đã trình bày 21 khối DDL, bảy vị trí chèn ảnh giao diện, bốn đoạn mã giao diện và kết quả xác minh. Việc nghiệm thu vận hành còn phụ thuộc vào quá trình migration runtime và smoke test sau migration.",
    )

    # Sửa mô tả trạng thái trong bảng yêu cầu chức năng.
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "Xử lý Request":
                row.cells[1].text = "Chuyển trạng thái Pending -> Accepted -> In_progress -> Completed hoặc Cancelled."
                break

    ensure_target_bookmarks(doc, manager)
    rebuild_front_matter(doc)
    add_missing_abbreviations(doc)
    reanchor_navigation_targets(doc, manager)

    # Bảo đảm màu chữ nội dung mới và phần note đã được loại bỏ hoàn toàn.
    remove_highlights(doc)
    doc.settings.update_fields_on_open = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
