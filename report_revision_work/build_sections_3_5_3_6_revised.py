from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

from build_sections_3_5_3_6 import add_body, add_bullet, set_cell_free_document_defaults, set_font


def build(output_path: Path) -> None:
    doc = Document()
    set_cell_free_document_defaults(doc)

    h35 = doc.add_paragraph(style="Heading 2")
    set_font(h35.add_run("3.5. Nhược điểm và hạn chế còn tồn tại"), size=14, bold=True)

    add_body(
        doc,
        "Mặc dù đã đáp ứng được các chức năng chính về cấu hình bàn phím, gửi yêu cầu lắp ráp, quản lý người dùng và kiểm tra chất lượng, phiên bản hiện tại vẫn tồn tại một số nhược điểm ảnh hưởng trực tiếp đến phạm vi sử dụng, trải nghiệm người dùng và khả năng triển khai thực tế.",
    )

    add_bullet(
        doc,
        "Chỉ sử dụng được trên máy tính Windows: ",
        "ứng dụng được xây dựng bằng WPF nên phải cài đặt và chạy trực tiếp trên hệ điều hành Windows. Hệ thống chưa thể sử dụng bằng trình duyệt, điện thoại hoặc máy tính dùng hệ điều hành khác, làm hạn chế khả năng tiếp cận của Buyer và Seller khi không sử dụng máy tính đã cài ứng dụng.",
    )
    add_bullet(
        doc,
        "Khả năng truy cập qua mạng còn hạn chế: ",
        "ứng dụng hiện kết nối trực tiếp tới SQL Server và chưa có lớp Web API trung gian. Việc sử dụng từ nhiều địa điểm phụ thuộc vào cấu hình mạng, chuỗi kết nối và quyền truy cập cơ sở dữ liệu. Kiến trúc này chưa phù hợp để công khai hệ thống trên Internet hoặc phục vụ số lượng lớn người dùng đồng thời.",
    )
    add_bullet(
        doc,
        "Chức năng chat chưa phải chat trực tiếp: ",
        "tin nhắn được lưu trong SQL Server nhưng SignalR hoặc WebSocket chưa được triển khai. Nội dung hội thoại có thể không xuất hiện ngay lập tức nếu giao diện chưa tải lại dữ liệu; các chức năng như trạng thái online, đang nhập, đã nhận, đã đọc, gửi tệp và thông báo tin nhắn mới cũng chưa có.",
    )
    add_bullet(
        doc,
        "Kiểm tra chất lượng mới ở mức mô phỏng: ",
        "DeviceSimulator tạo dữ liệu giả lập về phím, độ trễ và độ ồn để kiểm tra luồng nghiệp vụ. Do chưa kết nối ESP32, bộ kiểm tra ma trận phím hoặc cảm biến thật, kết quả QC hiện chỉ chứng minh logic phần mềm, chưa phản ánh chính xác chất lượng của một bàn phím vật lý.",
    )
    add_bullet(
        doc,
        "Quy trình đặt hàng chưa hoàn chỉnh: ",
        "hệ thống mới tập trung vào cấu hình build và gửi yêu cầu lắp ráp cho Seller. Các chức năng thường có trong một nền tảng thương mại điện tử như giỏ hàng, quản lý tồn kho, thanh toán trực tuyến, vận chuyển, theo dõi đơn hàng, hóa đơn và đánh giá sau mua chưa được triển khai.",
    )
    add_bullet(
        doc,
        "Khả năng vận hành và mở rộng chưa cao: ",
        "MQTT mới hoạt động theo cơ chế best-effort, bộ kiểm thử tự động chưa được khôi phục đầy đủ và migration cơ sở dữ liệu vận hành chưa được xác minh hoàn chỉnh. Hệ thống cũng chưa có cơ chế giám sát tập trung, cảnh báo lỗi, tự động sao lưu và quy trình triển khai liên tục cho môi trường thực tế.",
    )

    h36 = doc.add_paragraph(style="Heading 2")
    set_font(h36.add_run("3.6. Hướng phát triển trong tương lai"), size=14, bold=True)

    add_body(
        doc,
        "Trong giai đoạn tiếp theo, hệ thống có thể được phát triển theo hướng hoạt động trực tuyến, hỗ trợ trao đổi tức thời và hoàn thiện toàn bộ quy trình từ cấu hình bàn phím đến thanh toán, lắp ráp, kiểm tra và giao hàng.",
    )

    add_bullet(
        doc,
        "Phát triển chat trực tiếp theo thời gian thực: ",
        "SignalR hoặc WebSocket có thể được sử dụng để Buyer, Seller và Admin nhận tin nhắn ngay lập tức mà không cần tải lại giao diện. Chức năng chat nên bổ sung trạng thái online, đang nhập, đã nhận, đã đọc, thông báo tin nhắn mới, gửi hình ảnh hoặc tệp và lưu đầy đủ lịch sử hội thoại.",
    )
    add_bullet(
        doc,
        "Xây dựng phiên bản web hoạt động qua Internet: ",
        "các nghiệp vụ hiện có có thể được cung cấp thông qua ASP.NET Core Web API và một giao diện web responsive. Người dùng chỉ cần truy cập tên miền bằng trình duyệt để đăng nhập, cấu hình bàn phím, xử lý request hoặc quản trị hệ thống mà không phải cài đặt ứng dụng WPF trên từng máy.",
    )
    add_bullet(
        doc,
        "Hỗ trợ thiết bị di động: ",
        "giao diện web có thể được tối ưu cho điện thoại và máy tính bảng hoặc phát triển thành Progressive Web App. Buyer có thể theo dõi trạng thái lắp ráp, nhận thông báo và trao đổi với Seller ở bất kỳ đâu; Seller cũng có thể cập nhật tiến độ trực tiếp tại khu vực lắp ráp hoặc kiểm tra QC.",
    )
    add_bullet(
        doc,
        "Hoàn thiện quy trình mua bán trực tuyến: ",
        "hệ thống có thể bổ sung giỏ hàng, tồn kho, báo giá, mã giảm giá, thanh toán trực tuyến, lựa chọn đơn vị vận chuyển, theo dõi hành trình đơn hàng, xuất hóa đơn và đánh giá Seller. Khi đó, request lắp ráp sẽ trở thành một phần của quy trình đơn hàng hoàn chỉnh.",
    )
    add_bullet(
        doc,
        "Kết nối trạm QC và thiết bị IoT: ",
        "ESP32, bộ kiểm tra ma trận phím và cảm biến đo độ trễ hoặc độ ồn có thể được kết nối với hệ thống thông qua MQTT. Kết quả đo thực tế được gửi về máy chủ, lưu trong SQL Server và hiển thị trực tiếp trên dashboard để Seller và Admin theo dõi.",
    )
    add_bullet(
        doc,
        "Triển khai hệ thống trên hạ tầng trực tuyến: ",
        "Web API, cơ sở dữ liệu và MQTT broker có thể được triển khai trên máy chủ hoặc nền tảng điện toán đám mây với HTTPS, sao lưu định kỳ, quản lý bí mật, ghi log tập trung, health check và cảnh báo. Quy trình CI/CD giúp tự động build, kiểm thử và cập nhật phiên bản mới an toàn hơn.",
    )
    add_bullet(
        doc,
        "Bổ sung gợi ý cấu hình thông minh: ",
        "dựa trên ngân sách, layout, công nghệ switch, nhu cầu gõ và mức độ tiếng ồn mong muốn, hệ thống có thể đề xuất keyboard kit và linh kiện tương thích. Dữ liệu lịch sử cũng có thể được sử dụng để phân tích xu hướng, dự đoán nhu cầu và hỗ trợ Seller chuẩn bị linh kiện.",
    )

    add_body(
        doc,
        "Các hướng phát triển trên giúp khắc phục giới hạn của phiên bản desktop hiện tại, đồng thời chuyển hệ thống thành một nền tảng trực tuyến có khả năng phục vụ nhiều người dùng, hỗ trợ giao tiếp tức thời và kết nối với thiết bị kiểm tra phần cứng.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Nội dung bổ sung mục 3.5 và 3.6 - bản chỉnh sửa"
    doc.core_properties.subject = "Custom Keyboard Builder"
    doc.core_properties.author = ""
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: build_sections_3_5_3_6_revised.py OUTPUT.docx")
    build(Path(sys.argv[1]))
