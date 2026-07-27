from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FONT_NAME = "Times New Roman"


def set_font(run, size: float = 13, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "vi-VN")


def set_cell_free_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.07)
    section.footer_distance = Cm(1.07)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading = doc.styles["Heading 2"]
    heading.font.name = FONT_NAME
    heading.font.size = Pt(14)
    heading.font.bold = True
    heading.font.color.rgb = None
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.line_spacing = 1.15
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.widow_control = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = FONT_NAME
    bullet.font.size = Pt(13)
    bullet._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bullet.paragraph_format.left_indent = Cm(0.75)
    bullet.paragraph_format.first_line_indent = Cm(-0.45)
    bullet.paragraph_format.space_before = Pt(0)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.2
    bullet.paragraph_format.widow_control = True


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.keep_together = True
    set_font(paragraph.add_run(text))


def add_bullet(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.keep_together = True
    set_font(paragraph.add_run(label), bold=True)
    set_font(paragraph.add_run(text))


def build(output_path: Path) -> None:
    doc = Document()
    set_cell_free_document_defaults(doc)

    h35 = doc.add_paragraph(style="Heading 2")
    set_font(h35.add_run("3.5. Các lỗi và hạn chế còn tồn tại"), size=14, bold=True)

    add_body(
        doc,
        "Qua quá trình triển khai và xác minh, các chức năng cốt lõi của hệ thống đã có thành phần hiện thực tương ứng. Tuy nhiên, phiên bản hiện tại vẫn còn một số lỗi và giới hạn cần được xử lý trước khi có thể nghiệm thu ở mức vận hành ổn định.",
    )

    add_bullet(
        doc,
        "Cơ sở dữ liệu vận hành chưa được đồng bộ hoàn toàn: ",
        "mã nguồn, DBML và các script mới đã chuyển sang khóa chính id, nhưng cơ sở dữ liệu runtime được kiểm tra ngày 15/07/2026 vẫn sử dụng 21 khóa chính theo cấu trúc cũ. Migration PK-to-id đã được chuẩn bị nhưng chưa có đầy đủ bằng chứng về backup, diễn tập restore, postflight, rollback và smoke test sau migration.",
    )
    add_bullet(
        doc,
        "Bộ kiểm thử tự động chưa được khôi phục đầy đủ: ",
        "các runner Phase6Verification và WpfUiVerification không còn trong working tree nên chưa thể chạy lại toàn bộ kịch bản integration test và UI automation đã từng được sử dụng. Kết quả hiện tại chủ yếu dựa trên dotnet build, kiểm tra script và đối chiếu tĩnh, vì vậy mức độ bao phủ lỗi ở luồng giao diện vẫn còn hạn chế.",
    )
    add_bullet(
        doc,
        "Kiểm tra chất lượng mới dừng ở mức mô phỏng: ",
        "DeviceSimulator tạo dữ liệu kiểm tra phím, độ trễ và độ ồn để phục vụ luồng QC. Do chưa kết nối ESP32 hoặc trạm đo phần cứng thật, các kết quả pass, warning và fail hiện chỉ phản ánh logic phần mềm, chưa thể được xem là kết quả kiểm định vật lý của bàn phím.",
    )
    add_bullet(
        doc,
        "Khả năng truyền thông thời gian thực chưa hoàn thiện: ",
        "MQTT được triển khai theo cơ chế best-effort và ứng dụng chuyển sang chế độ DB-only khi broker không khả dụng. Chức năng chat hiện lưu dữ liệu trong SQL Server nhưng SignalR chưa được triển khai, vì vậy khả năng nhận tin tức thời, trạng thái đã đọc và đồng bộ hội thoại giữa nhiều phiên làm việc chưa được bảo đảm đầy đủ.",
    )
    add_bullet(
        doc,
        "Một số ràng buộc dữ liệu chưa được thể hiện đầy đủ ở tầng cơ sở dữ liệu: ",
        "các cột mang ý nghĩa giá trị liệt kê, tiêu biểu như role_name và switch_technology, chưa có CHECK constraint hoàn chỉnh. Dữ liệu được Service kiểm tra trước khi ghi, nhưng giá trị không hợp lệ vẫn có thể xuất hiện nếu dữ liệu được thêm trực tiếp bằng script hoặc công cụ quản trị.",
    )
    add_bullet(
        doc,
        "Xử lý đồng thời và chuyển trạng thái cần được kiểm thử thêm: ",
        "request đã có luồng Pending, Accepted, In_progress, Completed hoặc Cancelled, nhưng cơ chế versioning, hủy thao tác và kiểm soát cập nhật đồng thời chưa được xác minh đầy đủ. Khi nhiều thao tác cùng cập nhật một request, nguy cơ ghi đè trạng thái hoặc sử dụng dữ liệu cũ vẫn cần được loại trừ.",
    )

    h36 = doc.add_paragraph(style="Heading 2")
    set_font(h36.add_run("3.6. Hướng phát triển trong tương lai"), size=14, bold=True)

    add_body(
        doc,
        "Các hướng phát triển tiếp theo được xác định theo thứ tự ưu tiên từ ổn định dữ liệu, tăng độ tin cậy của kiểm thử đến mở rộng khả năng kết nối và phạm vi nghiệp vụ của hệ thống.",
    )

    add_bullet(
        doc,
        "Hoàn tất migration cơ sở dữ liệu: ",
        "cần tạo bản backup, diễn tập restore trên môi trường tách biệt, chạy migration PK-to-id, kiểm tra postflight và chuẩn bị phương án rollback. Sau migration, toàn bộ luồng đăng nhập, cấu hình build, xử lý request, chat, audit và QC cần được smoke test trước khi sử dụng cơ sở dữ liệu mới.",
    )
    add_bullet(
        doc,
        "Khôi phục và mở rộng kiểm thử tự động: ",
        "các runner integration/UI cần được đưa trở lại dự án và bổ sung vào quy trình build tự động. Bộ kiểm thử nên bao phủ phân quyền Buyer, Seller và Admin; tính tương thích linh kiện; state machine của request; giao dịch cơ sở dữ liệu; các trường hợp MQTT không khả dụng và những lỗi nhập liệu thường gặp.",
    )
    add_bullet(
        doc,
        "Tăng cường tính toàn vẹn và an toàn dữ liệu: ",
        "cần bổ sung CHECK constraint cho các giá trị liệt kê, kiểm thử SQL injection, dữ liệu biên và phân quyền truy cập. Cơ chế optimistic concurrency hoặc versioning nên được áp dụng cho request, đồng thời hoàn thiện quy trình cancellation để tránh ghi đè trạng thái khi có nhiều thao tác đồng thời.",
    )
    add_bullet(
        doc,
        "Hoàn thiện chức năng realtime: ",
        "SignalR có thể được bổ sung cho chat để hỗ trợ nhận tin tức thời, trạng thái đã đọc và thông báo hội thoại mới. MQTT tiếp tục được sử dụng cho request, trạng thái và telemetry QC, nhưng cần bổ sung cơ chế reconnect, hàng đợi tạm và theo dõi lỗi broker để tăng độ ổn định.",
    )
    add_bullet(
        doc,
        "Kết nối trạm QC phần cứng: ",
        "DeviceSimulator có thể được thay thế hoặc kết hợp với ESP32, bộ kiểm tra ma trận phím và cảm biến đo độ trễ hoặc độ ồn. Dữ liệu đo thực tế sẽ được truyền qua MQTT, lưu vào SQL Server và đối chiếu với ngưỡng kiểm tra để tạo kết quả QC có giá trị thực nghiệm.",
    )
    add_bullet(
        doc,
        "Mở rộng phạm vi nghiệp vụ và vận hành: ",
        "sau khi các chức năng cốt lõi ổn định, hệ thống có thể bổ sung quản lý tồn kho, thanh toán, vận chuyển, theo dõi đơn hàng và thông báo đa kênh. Đồng thời, cần xây dựng lịch backup, nhật ký tập trung, health check và tài liệu triển khai để hỗ trợ vận hành lâu dài.",
    )

    add_body(
        doc,
        "Việc triển khai các hướng trên sẽ giúp hệ thống chuyển từ mô hình minh họa nghiệp vụ sang một nền tảng có khả năng vận hành ổn định, kiểm thử lặp lại và mở rộng sang thiết bị phần cứng trong tương lai.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Nội dung bổ sung mục 3.5 và 3.6"
    doc.core_properties.subject = "Custom Keyboard Builder"
    doc.core_properties.author = ""
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: build_sections_3_5_3_6.py OUTPUT.docx")
    build(Path(sys.argv[1]))
