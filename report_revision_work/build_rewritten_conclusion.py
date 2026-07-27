from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from build_sections_3_5_3_6 import FONT_NAME, add_body, set_cell_free_document_defaults, set_font


def configure_conclusion_title(doc: Document) -> None:
    style = doc.styles["Heading 1"]
    style.font.name = FONT_NAME
    style.font.size = Pt(16)
    style.font.bold = True
    style.font.color.rgb = None
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(10)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.keep_with_next = True


def add_heading_2(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 2")
    set_font(paragraph.add_run(text), size=14, bold=True)


def build(output_path: Path) -> None:
    doc = Document()
    set_cell_free_document_defaults(doc)
    configure_conclusion_title(doc)

    title = doc.add_paragraph(style="Heading 1")
    set_font(title.add_run("KẾT LUẬN"), size=16, bold=True)

    add_body(
        doc,
        "Đề tài Custom Keyboard Builder được thực hiện với mục tiêu xây dựng một hệ thống hỗ trợ cấu hình, quản lý và gửi yêu cầu lắp ráp bàn phím tùy chỉnh. Quá trình thực hiện đã bao gồm khảo sát bài toán, xác định yêu cầu, phân tích thiết kế, xây dựng cơ sở dữ liệu, triển khai ứng dụng và kiểm tra các luồng nghiệp vụ chính.",
    )

    add_heading_2(doc, "1. Mức độ hoàn thành nhiệm vụ")
    add_body(
        doc,
        "Các nhiệm vụ phân tích và thiết kế đã được hoàn thành thông qua việc xác định ba tác nhân Buyer, Seller và Admin; xây dựng yêu cầu chức năng, yêu cầu phi chức năng và năm sơ đồ gồm Use Case, Activity, Sequence, ERD tổng quát và ERD chi tiết. Cơ sở dữ liệu SQL Server được thiết kế theo 21 bảng nghiệp vụ, thể hiện các nhóm dữ liệu về tài khoản, danh mục linh kiện, cấu hình build, request, kiểm tra QC, chat và audit log.",
    )
    add_body(
        doc,
        "Ở mức triển khai, ứng dụng desktop WPF trên nền tảng .NET 10 đã được tổ chức theo MVVM kết hợp Service, Repository và Data/Realtime. Các chức năng cốt lõi đã có thành phần hiện thực tương ứng, bao gồm đăng ký và đăng nhập, phân quyền theo vai trò, lựa chọn linh kiện, lưu cấu hình, gửi và xử lý yêu cầu lắp ráp, quản lý Seller, chat lưu trong cơ sở dữ liệu, thống kê dashboard, ghi lịch sử thao tác và kiểm tra chất lượng bằng DeviceSimulator.",
    )
    add_body(
        doc,
        "Từ các kết quả trên, nhiệm vụ có thể được đánh giá là đã hoàn thành ở mức xây dựng mô hình, mã nguồn và minh họa được các luồng nghiệp vụ chính. Tuy nhiên, hệ thống chưa đạt mức sản phẩm vận hành hoàn chỉnh do cơ sở dữ liệu runtime chưa hoàn tất xác minh migration, bộ kiểm thử tự động chưa được khôi phục đầy đủ và một số chức năng vẫn dừng ở mức mô phỏng hoặc thử nghiệm.",
    )

    add_heading_2(doc, "2. Mức độ phù hợp với mục đích đề tài")
    add_body(
        doc,
        "Kết quả đạt được phù hợp với mục đích số hóa quy trình cấu hình và đặt lắp ráp bàn phím tùy chỉnh. Mô hình kit-based giúp giảm độ phức tạp so với việc lựa chọn rời từng bộ phận; các quy tắc tương thích, snapshot giá và cơ chế phân quyền góp phần hạn chế cấu hình không hợp lệ và duy trì tính nhất quán của dữ liệu.",
    )
    add_body(
        doc,
        "Việc lưu trữ tập trung trên SQL Server giúp Buyer theo dõi cấu hình và trạng thái request, Seller quản lý quá trình tiếp nhận, lắp ráp và QC, còn Admin có thể quản trị người dùng, danh mục và lịch sử hoạt động. Vì vậy, hệ thống đã giải quyết được phần chính của bài toán đặt ra và tạo nền tảng để tiếp tục phát triển thành một dịch vụ trực tuyến hoàn chỉnh.",
    )

    add_heading_2(doc, "3. Những hạn chế chính")
    add_body(
        doc,
        "Phiên bản hiện tại chỉ hoạt động dưới dạng ứng dụng WPF trên Windows, phải cài đặt trên từng máy và chưa thể truy cập thuận tiện bằng trình duyệt hoặc thiết bị di động. Ứng dụng kết nối trực tiếp tới SQL Server, chưa có Web API trung gian nên khả năng triển khai qua Internet, phục vụ nhiều người dùng đồng thời và mở rộng trên hạ tầng đám mây còn hạn chế.",
    )
    add_body(
        doc,
        "Chức năng chat mới lưu tin nhắn trong cơ sở dữ liệu, chưa hỗ trợ trao đổi tức thời bằng SignalR hoặc WebSocket. Quy trình QC sử dụng dữ liệu mô phỏng nên chưa phản ánh chất lượng bàn phím vật lý. Ngoài ra, các chức năng giỏ hàng, tồn kho, thanh toán, vận chuyển, theo dõi đơn hàng, giám sát hệ thống và kiểm thử tự động đầy đủ vẫn chưa được hoàn thiện.",
    )

    add_heading_2(doc, "4. Hướng phát triển")
    add_body(
        doc,
        "Trong tương lai, chức năng chat có thể được phát triển theo thời gian thực để hỗ trợ trạng thái online, đang nhập, đã nhận, đã đọc, gửi hình ảnh hoặc tệp và thông báo tin nhắn mới. Phiên bản web hoạt động qua Internet có thể được xây dựng bằng ASP.NET Core Web API kết hợp giao diện responsive, cho phép Buyer, Seller và Admin sử dụng hệ thống trên trình duyệt, điện thoại hoặc máy tính bảng mà không cần cài đặt ứng dụng WPF.",
    )
    add_body(
        doc,
        "Quy trình nghiệp vụ có thể được mở rộng với quản lý tồn kho, giỏ hàng, thanh toán trực tuyến, vận chuyển và theo dõi đơn hàng. DeviceSimulator có thể được kết hợp hoặc thay thế bằng ESP32, bộ kiểm tra ma trận phím và cảm biến thực tế; dữ liệu QC được truyền qua MQTT và hiển thị trực tiếp trên dashboard. Đồng thời, việc triển khai trên máy chủ hoặc nền tảng điện toán đám mây, bổ sung HTTPS, sao lưu, logging, health check, cảnh báo và CI/CD sẽ giúp hệ thống hoạt động ổn định hơn.",
    )

    add_body(
        doc,
        "Nhìn chung, đề tài đã hình thành được một nền tảng phần mềm phù hợp với mục tiêu quản lý cấu hình và yêu cầu lắp ráp bàn phím tùy chỉnh. Kết quả hiện tại đáp ứng yêu cầu học tập, thể hiện được quá trình phân tích, thiết kế và triển khai hệ thống; đồng thời để lại hướng mở rõ ràng để phát triển thành một nền tảng trực tuyến có khả năng giao tiếp thời gian thực, phục vụ nhiều người dùng và kết nối với thiết bị kiểm tra phần cứng.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Kết luận báo cáo Custom Keyboard Builder"
    doc.core_properties.subject = "Custom Keyboard Builder"
    doc.core_properties.author = ""
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: build_rewritten_conclusion.py OUTPUT.docx")
    build(Path(sys.argv[1]))
