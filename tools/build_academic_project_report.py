from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = ROOT / "Database" / "SqlServer" / "CreateSchema_Refactor.sql"
LOGO_PATH = ROOT / "report_assets" / "hust_logo.jpg"
LOGIN_IMAGE = ROOT / "report_assets" / "ui_login_mock.png"
BUYER_IMAGE = ROOT / "report_assets" / "ui_buyer_dashboard_mock.png"
DEFAULT_OUTPUT = ROOT / "output" / "Bao_cao_du_an_Custom_Keyboard.docx"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

TABLE_ORDER = [
    "roles",
    "users",
    "seller_profiles",
    "seller_applications",
    "brands",
    "layouts",
    "keyboard_kits",
    "switches",
    "keycap_sets",
    "stabilizers",
    "accessories",
    "builds",
    "build_items",
    "build_mods",
    "build_requests",
    "devices",
    "device_test_sessions",
    "device_key_test_results",
    "audit_log",
    "chat_conversations",
    "chat_messages",
]


@dataclass(frozen=True)
class TocEntry:
    anchor: str
    level: int
    text: str


@dataclass(frozen=True)
class FigureDef:
    anchor: str
    number: str
    title: str
    group: str
    source: str
    description: str
    blank_height_in: float = 3.1

    @property
    def label(self) -> str:
        return f"Hình {self.number}. {self.title}"


TOC_ENTRIES = [
    TocEntry("opening", 1, "MỞ ĐẦU"),
    TocEntry("opening_1", 2, "1. Giới thiệu và lý do chọn đề tài"),
    TocEntry("opening_2", 2, "2. Mục đích và nhiệm vụ của đề tài"),
    TocEntry("opening_3", 2, "3. Phân công nhiệm vụ của nhóm"),
    TocEntry("opening_4", 2, "4. Nội dung sơ lược của báo cáo"),
    TocEntry("ch1", 1, "CHƯƠNG 1. TỔNG QUAN HỆ THỐNG CUSTOM KEYBOARD BUILDER"),
    TocEntry("ch1_1", 2, "1.1. Bối cảnh và vấn đề cần giải quyết"),
    TocEntry("ch1_data", 2, "1.2. Yêu cầu về dữ liệu"),
    TocEntry("ch1_2", 2, "1.3. Mục tiêu và phạm vi hệ thống"),
    TocEntry("ch1_3", 2, "1.4. Tác nhân và yêu cầu chức năng"),
    TocEntry("ch1_4", 2, "1.5. Yêu cầu phi chức năng"),
    TocEntry("ch1_5", 2, "1.6. Quy tắc nghiệp vụ và ranh giới đề tài"),
    TocEntry("ch1_6", 2, "1.7. Kết luận chương"),
    TocEntry("ch2", 1, "CHƯƠNG 2. PHÂN TÍCH THIẾT KẾ HỆ THỐNG"),
    TocEntry("ch2_1", 2, "2.1. Công nghệ và môi trường triển khai"),
    TocEntry("ch2_2", 2, "2.2. Kiến trúc phần mềm"),
    TocEntry("ch2_3", 2, "2.3. Entity-Relationship Diagram (ERD)"),
    TocEntry("ch2_4", 2, "2.4. Use Case Diagram"),
    TocEntry("ch2_5", 2, "2.5. Sequence Diagram"),
    TocEntry("ch2_6", 2, "2.6. Activity Diagram"),
    TocEntry("ch2_7", 2, "2.7. Đối chiếu diagram với thành phần triển khai"),
    TocEntry("ch2_8", 2, "2.8. Kết luận chương"),
    TocEntry("ch3", 1, "CHƯƠNG 3. TRIỂN KHAI VÀ KIỂM THỬ"),
    TocEntry("ch3_1", 2, "3.1. Triển khai cơ sở dữ liệu SQL Server"),
    TocEntry("ch3_1_1", 3, "3.1.1. Nhóm tài khoản và phân quyền"),
    TocEntry("ch3_1_2", 3, "3.1.2. Nhóm danh mục linh kiện"),
    TocEntry("ch3_1_3", 3, "3.1.3. Nhóm build và request"),
    TocEntry("ch3_1_4", 3, "3.1.4. Nhóm thiết bị và kiểm tra QC"),
    TocEntry("ch3_1_5", 3, "3.1.5. Nhóm audit và chat"),
    TocEntry("ch3_2", 2, "3.2. Triển khai giao diện tiêu biểu"),
    TocEntry("ch3_3", 2, "3.3. Kiểm thử và xác minh"),
    TocEntry("ch3_4", 2, "3.4. Đánh giá mức độ hoàn thành"),
    TocEntry("ch3_5", 2, "3.5. Kết luận chương"),
    TocEntry("conclusion", 1, "KẾT LUẬN"),
    TocEntry("conclusion_1", 2, "1. Mức độ hoàn thành nhiệm vụ"),
    TocEntry("conclusion_2", 2, "2. Mức độ phù hợp với mục đích đề tài"),
    TocEntry("conclusion_3", 2, "3. Hạn chế còn tồn tại"),
    TocEntry("conclusion_4", 2, "4. Hướng phát triển"),
]


FIGURES = [
    FigureDef(
        "fig_2_1", "2.1", "Overall Entity-Relationship Diagram of the Custom Keyboard Builder System", "ERD",
        "Documents_Refactor/Custom_Keyboard_ERD_Realistic_Kit_Shop_Proposal.dbml",
        "Sơ đồ cần thể hiện các vùng dữ liệu tài khoản, catalog, build/request, QC, audit và chat cùng quan hệ chính giữa các vùng.",
    ),
    FigureDef(
        "fig_2_2", "2.2", "Detailed Entity-Relationship Diagram with 21 Tables and 33 Relationships", "ERD",
        "Documents_Refactor/Custom_Keyboard_ERD_Realistic_Kit_Shop_Proposal.dbml",
        "Sơ đồ chi tiết cần hiển thị cột, kiểu dữ liệu, khóa chính id, khóa ngoại và các ràng buộc quan trọng theo DBML.",
    ),
    FigureDef(
        "fig_2_3", "2.3", "Use Case UC-01 - Buyer Creates a Build and Sends a Request", "Use Case",
        "Documents_Refactor/KTPMUD_diagram/16_UseCase_UC01_Buyer.drawio",
        "Sơ đồ thể hiện các ca sử dụng đăng ký, cấu hình build, lưu build, chọn Seller, gửi request, theo dõi QC và chat.",
    ),
    FigureDef(
        "fig_2_4", "2.4", "Use Case UC-02 - Seller Processes a Request and Performs QC", "Use Case",
        "Documents_Refactor/KTPMUD_diagram/17_UseCase_UC02_Seller.drawio",
        "Sơ đồ thể hiện các ca sử dụng xem request, chuyển trạng thái, chạy QC, xem kết quả và chat với Buyer/Admin.",
    ),
    FigureDef(
        "fig_2_5", "2.5", "Use Case UC-03 - Admin Manages the System", "Use Case",
        "Documents_Refactor/KTPMUD_diagram/18_UseCase_UC03_Admin.drawio",
        "Sơ đồ thể hiện các ca sử dụng quản lý user, Seller, catalog, đơn đăng ký Seller, audit log và chat với Seller.",
    ),
    FigureDef(
        "fig_2_6", "2.6", "SD-R01 - Registration, Login, and Role-Based Navigation", "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ mô tả trình tự giữa giao diện, AccountService, repository và SQL Server trong xác thực và mở dashboard.",
    ),
    FigureDef(
        "fig_2_7", "2.7", "SD-R02 - Buyer Creates a Build and Sends a Request", "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ mô tả trình tự lấy catalog, validate, lưu build, tạo snapshot request và phát sự kiện realtime best-effort.",
    ),
    FigureDef(
        "fig_2_8", "2.8", "SD-R03 - Seller Processes a Request and QC Results", "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ mô tả Seller đọc request, cập nhật trạng thái, chạy DeviceSimulator và nhận QC summary.",
    ),
    FigureDef(
        "fig_2_9", "2.9", "SD-R04 - Admin Manages the System", "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ mô tả trình tự tải dashboard, quản lý user/catalog/Seller và ghi audit log.",
    ),
    FigureDef(
        "fig_2_10", "2.10", "SD-R05 - Seller Application Approval", "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ mô tả Buyer nộp đơn, Admin duyệt, đổi role và tạo Seller profile đã xác minh trong một giao dịch nghiệp vụ.",
    ),
    FigureDef(
        "fig_2_11", "2.11", "SD-R06 - Chat Between Authorized Users", "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ mô tả tạo hội thoại, kiểm tra participant, lưu tin nhắn và tải lại lịch sử chat.",
    ),
    FigureDef(
        "fig_2_12", "2.12", "AD-01 - Account Activities", "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ mô tả rẽ nhánh đăng ký/đăng nhập, kiểm tra dữ liệu, xem hồ sơ và đăng xuất.",
    ),
    FigureDef(
        "fig_2_13", "2.13", "AD-02 - Buyer Creates and Submits a Build", "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ mô tả chuỗi hoạt động chọn kit, thêm linh kiện, validate, lưu build, chọn Seller và gửi request.",
    ),
    FigureDef(
        "fig_2_14", "2.14", "AD-03 - Seller Processes a Request", "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ mô tả các nhánh chấp nhận, đang xử lý, chạy QC, hoàn thành hoặc hủy request.",
    ),
    FigureDef(
        "fig_2_15", "2.15", "AD-04 - Admin Management Activities", "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ mô tả các hoạt động quản lý user, Seller, catalog, duyệt đơn Seller và xem audit log.",
    ),
    FigureDef(
        "fig_2_16", "2.16", "AD-05 - Buyer-Seller and Seller-Admin Chat", "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ mô tả chọn hội thoại, kiểm tra quyền, gửi/lưu tin nhắn và tải lịch sử trao đổi.",
    ),
    FigureDef(
        "fig_2_17", "2.17", "AD-06 - Keyboard Device/QC Testing", "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ mô tả tạo phiên QC, sinh telemetry từng phím, đánh giá pass/warning/fail và tổng hợp kết quả.",
    ),
    FigureDef(
        "fig_3_1", "3.1", "Giao diện đăng nhập của ứng dụng", "UI",
        "Views/LoginView.xaml",
        "Giao diện gồm vùng giới thiệu, form đăng nhập, thông báo trạng thái và điều hướng tạo tài khoản Buyer.",
    ),
    FigureDef(
        "fig_3_2", "3.2", "Giao diện Buyer Dashboard và khu vực cấu hình build", "UI",
        "Views/BuyerDashboardView.xaml",
        "Giao diện thể hiện sidebar, configurator kit-based, tổng giá snapshot, validation và chức năng gửi request.",
    ),
]


TABLES = {
    "tbl_0_1": ("0.1", "Phân công nhiệm vụ nhóm"),
    "tbl_1_data": ("1.1", "Yêu cầu về dữ liệu"),
    "tbl_1_1": ("1.2", "Tác nhân và mục tiêu sử dụng hệ thống"),
    "tbl_1_2": ("1.3", "Yêu cầu chức năng chính"),
    "tbl_1_3": ("1.4", "Yêu cầu phi chức năng"),
    "tbl_1_4": ("1.5", "Phạm vi thực hiện và nội dung ngoài phạm vi"),
    "tbl_2_1": ("2.1", "Công nghệ và thư viện sử dụng"),
    "tbl_2_2": ("2.2", "Các lớp trong kiến trúc phần mềm"),
    "tbl_2_3": ("2.3", "Nhóm bảng trong cơ sở dữ liệu"),
    "tbl_2_4": ("2.4", "Đối chiếu bốn loại diagram với artifact nguồn"),
    "tbl_3_1": ("3.1", "Thứ tự triển khai 21 bảng theo nhóm phụ thuộc"),
    "tbl_3_2": ("3.2", "Đối chiếu giao diện với mã nguồn"),
    "tbl_3_3": ("3.3", "Kết quả build và xác minh tại thời điểm lập báo cáo"),
    "tbl_3_4": ("3.4", "Đánh giá mức độ hoàn thành theo mục tiêu"),
}


def set_run_font(run, name: str = "Times New Roman", size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 color: str | None = None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_inline_picture_alt(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def set_cell_margins(cell, top: int = 90, start: int = 120,
                     bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, attrs in (
        ("w:tblW", {"w:w": str(CONTENT_WIDTH_DXA), "w:type": "dxa"}),
        ("w:tblInd", {"w:w": str(indent_dxa), "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ):
        node = tbl_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tbl_pr.append(node)
        for key, value in attrs.items():
            node.set(qn(key), value)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])


def set_table_borders(table, color: str = "A6A6A6", size: str = "6",
                      style: str = "single") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), style)
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    set_table_borders(table, color="FFFFFF", size="0", style="nil")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "true")
        tr_pr.append(marker)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:cantSplit"))
    if marker is None:
        marker = OxmlElement("w:cantSplit")
        tr_pr.append(marker)


def set_page_geometry(section, cover: bool = False) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    if cover:
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.62)
        section.right_margin = Inches(0.62)
    else:
        # Match the sample's academic layout while preserving a 6.5-inch text block.
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)


def add_page_border(section) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgBorders"))
    if existing is not None:
        sect_pr.remove(existing)
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "double")
        node.set(qn("w:sz"), "12")
        node.set(qn("w:space"), "18")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    sect_pr.append(borders)


def remove_page_border(section) -> None:
    existing = section._sectPr.find(qn("w:pgBorders"))
    if existing is not None:
        section._sectPr.remove(existing)


def append_field(paragraph, instruction: str, cached_text: str = "1") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:r")
    result_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{attr}"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    result_pr.extend([fonts, size])
    result.append(result_pr)
    text = OxmlElement("w:t")
    text.text = cached_text
    result.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    # Field characters and instructions must live inside runs; placing them as
    # direct paragraph children produces a package that Word refuses to open.
    begin_run = OxmlElement("w:r")
    begin_run.append(begin)
    instr_run = OxmlElement("w:r")
    instr_run.append(instr)
    separate_run = OxmlElement("w:r")
    separate_run.append(separate)
    end_run = OxmlElement("w:r")
    end_run.append(end)
    paragraph._p.extend([begin_run, instr_run, separate_run, result, end_run])


def set_page_footer(section) -> None:
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    append_field(paragraph, " PAGE ", "1")


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, 14, 8, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": (14, 10, 5, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 3": (13, 8, 4, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for style_name, (size, before, after, alignment) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
            r_fonts.set(qn(f"w:{attr}"), "Times New Roman")
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Front Matter Title" not in styles:
        style = styles.add_style("Front Matter Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Front Matter Title"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(16)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(18)
    style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        if list_name not in styles:
            continue
        style = styles[list_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(13)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2


class BookmarkManager:
    def __init__(self) -> None:
        self._next_id = 1

    def add(self, paragraph, name: str) -> None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._next_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._next_id))
        paragraph._p.insert(1, start)
        paragraph._p.append(end)
        self._next_id += 1


def add_internal_hyperlink(paragraph, text: str, anchor: str, size: float = 13,
                           bold: bool = False, color: str = "0563C1") -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), "Times New Roman")
    r_pr.append(fonts)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_node)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    if bold:
        bold_node = OxmlElement("w:b")
        r_pr.append(bold_node)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc: Document, manager: BookmarkManager, anchor: str, text: str,
                level: int, page_break_before: bool | None = None):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if page_break_before is None:
        page_break_before = level == 1
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 14 if level == 2 else 13, bold=True)
    manager.add(paragraph, anchor)
    return paragraph


def add_body(doc: Document, text: str, keep_next: bool = False, italic: bool = False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, size: float = 13):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = align
    paragraph.paragraph_format.keep_with_next = keep_next
    run = paragraph.add_run(text)
    set_run_font(run, size=size, italic=italic)
    return paragraph


def add_linked_body(doc: Document, segments: list[tuple[str, str | None]],
                    keep_next: bool = False, size: float = 13):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.keep_with_next = keep_next
    for text, anchor in segments:
        if anchor is None:
            run = paragraph.add_run(text)
            set_run_font(run, size=size)
        else:
            add_internal_hyperlink(paragraph, text, anchor, size=size)
    return paragraph


def add_bullet(doc: Document, text: str, level: int = 1):
    style = "List Bullet" if level == 1 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, size=13)
    return paragraph


def add_front_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Front Matter Title")
    run = paragraph.add_run(text)
    set_run_font(run, size=16, bold=True)


def add_nav_line(doc: Document, text: str, anchor: str, page: str,
                 level: int = 1, size: float = 12, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25 * (level - 1))
    paragraph.paragraph_format.first_line_indent = Inches(-0.08 if level > 1 else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.35), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    add_internal_hyperlink(paragraph, text, anchor, size=size, bold=bold, color="000000")
    run = paragraph.add_run(f"\t{page}")
    set_run_font(run, size=size, bold=bold)


def add_table_caption(doc: Document, manager: BookmarkManager, anchor: str):
    number, title = TABLES[anchor]
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"Bảng {number}. {title}")
    set_run_font(run, size=12, bold=True)
    manager.add(paragraph, anchor)
    return paragraph


def add_data_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int],
                   header_fill: str = "E7E6E6"):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table, color="8C8C8C", size="6")
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, top=95, bottom=95)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(header)
        set_run_font(run, size=10.5, bold=True)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            set_cell_margins(cell, top=85, bottom=85)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_blank_diagram(doc: Document, manager: BookmarkManager, figure: FigureDef) -> None:
    add_linked_body(
        doc,
        [
            (figure.description + " Vị trí chèn thủ công được đánh dấu tại ", None),
            (f"Hình {figure.number}", figure.anchor),
            (".", None),
        ],
        keep_next=True,
    )
    source = add_body(doc, f"Nguồn diagram: {figure.source}", keep_next=True, italic=True, size=10.5)
    source.paragraph_format.space_after = Pt(4)
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA], indent_dxa=0)
    set_table_borders(table, color="BFBFBF", size="5", style="dashed")
    row = table.rows[0]
    # Named override: an empty exact-height row deliberately reserves replaceable figure space.
    row.height = Inches(figure.blank_height_in)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True
    run = caption.add_run(figure.label)
    set_run_font(run, size=12, bold=True)
    manager.add(caption, figure.anchor)


def add_picture_figure(doc: Document, manager: BookmarkManager, figure: FigureDef,
                       image_path: Path) -> None:
    add_linked_body(
        doc,
        [
            (figure.description + " Bố cục được minh họa tại ", None),
            (f"Hình {figure.number}", figure.anchor),
            (".", None),
        ],
        keep_next=True,
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(6.15))
    set_inline_picture_alt(picture, figure.label)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run(figure.label)
    set_run_font(run, size=12, bold=True)
    manager.add(caption, figure.anchor)


SQL_KEYWORDS = {
    "CREATE", "TABLE", "CONSTRAINT", "PRIMARY", "KEY", "FOREIGN", "REFERENCES",
    "CHECK", "DEFAULT", "NOT", "NULL", "UNIQUE", "IDENTITY", "IN", "OR", "WHEN",
    "THEN", "ELSE", "END", "CASE", "INDEX", "ON", "WHERE", "DESC",
}
SQL_TYPES = {"INT", "BIGINT", "BIT", "VARCHAR", "NVARCHAR", "DATETIME2", "DECIMAL", "MAX"}
SQL_TOKEN_RE = re.compile(
    r"(--.*$|'(?:''|[^'])*'|\b[A-Za-z_][A-Za-z0-9_]*\b)", re.MULTILINE
)


def add_code_runs(paragraph, code: str, language: str) -> None:
    lines = code.strip("\n").splitlines()
    for line_index, line in enumerate(lines):
        if line_index:
            paragraph.add_run().add_break()
        if language.lower() != "sql":
            run = paragraph.add_run(line.rstrip())
            set_run_font(run, name="Consolas", size=7.6, color="1F2937")
            continue
        position = 0
        for match in SQL_TOKEN_RE.finditer(line.rstrip()):
            if match.start() > position:
                run = paragraph.add_run(line[position:match.start()])
                set_run_font(run, name="Consolas", size=7.4, color="1F2937")
            token = match.group(0)
            upper = token.upper()
            if token.startswith("--"):
                color = "6B7280"
            elif token.startswith("'"):
                color = "A31515"
            elif upper in SQL_KEYWORDS:
                color = "0000FF"
            elif upper in SQL_TYPES:
                color = "2B91AF"
            else:
                color = "1F2937"
            run = paragraph.add_run(token)
            set_run_font(run, name="Consolas", size=7.4, color=color)
            position = match.end()
        if position < len(line.rstrip()):
            run = paragraph.add_run(line[position:].rstrip())
            set_run_font(run, name="Consolas", size=7.4, color="1F2937")


def add_code_table(doc: Document, title: str, code: str, language: str) -> None:
    table = doc.add_table(rows=2, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA], indent_dxa=130)
    set_table_borders(table, color="9CA3AF", size="6")
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    prevent_row_split(table.rows[1])
    title_cell = table.rows[0].cells[0]
    code_cell = table.rows[1].cells[0]
    set_cell_shading(title_cell, "D9D9D9")
    set_cell_shading(code_cell, "FFFFFF")
    for cell in (title_cell, code_cell):
        set_cell_margins(cell, top=85, bottom=85, start=130, end=130)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    title_paragraph = title_cell.paragraphs[0]
    title_paragraph.paragraph_format.space_after = Pt(0)
    title_paragraph.paragraph_format.line_spacing = 1.0
    title_paragraph.paragraph_format.keep_with_next = True
    run = title_paragraph.add_run(title)
    set_run_font(run, size=10.5, bold=True)
    code_paragraph = code_cell.paragraphs[0]
    code_paragraph.paragraph_format.space_before = Pt(0)
    code_paragraph.paragraph_format.space_after = Pt(0)
    code_paragraph.paragraph_format.line_spacing = 1.0
    code_paragraph.paragraph_format.keep_together = True
    code_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_code_runs(code_paragraph, code, language)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def strip_sql_comments(block: str) -> str:
    cleaned_lines: list[str] = []
    for raw_line in block.splitlines():
        line = re.sub(r"\s+--.*$", "", raw_line).rstrip()
        if line.lstrip().startswith("--"):
            continue
        cleaned_lines.append(line)
    while cleaned_lines and not cleaned_lines[-1].strip():
        cleaned_lines.pop()
    return "\n".join(cleaned_lines)


def extract_sql_blocks() -> dict[str, str]:
    sql = SCHEMA_PATH.read_text(encoding="utf-8")
    pattern = re.compile(
        r"(?ms)^CREATE\s+TABLE\s+(?P<name>[a-z_][a-z0-9_]*)\s*\(.*?^\);\s*(?=^GO\s*$)"
    )
    blocks: dict[str, str] = {}
    for match in pattern.finditer(sql):
        name = match.group("name")
        if name in TABLE_ORDER:
            blocks[name] = strip_sql_comments(match.group(0))
    missing = [name for name in TABLE_ORDER if name not in blocks]
    if missing:
        raise RuntimeError(f"Missing CREATE TABLE blocks: {missing}")
    blocks["seller_applications"] += (
        "\n\nCREATE UNIQUE INDEX UX_seller_applications_pending_buyer\n"
        "ON seller_applications(buyer_user_id)\n"
        "WHERE status = 'Pending';"
    )
    blocks["device_key_test_results"] += (
        "\n\nCREATE UNIQUE INDEX UX_dktr_session_key_code\n"
        "ON device_key_test_results(session_id, key_code);"
    )
    return blocks


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    set_page_geometry(section, cover=True)
    add_page_border(section)
    set_page_footer(section)

    def centered(text: str, size: float, bold: bool = False, after: float = 6) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)

    centered("ĐẠI HỌC BÁCH KHOA HÀ NỘI", 16, True, 5)
    centered("TRƯỜNG ĐIỆN - ĐIỆN TỬ", 16, True, 8)
    if LOGO_PATH.exists():
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.paragraph_format.space_after = Pt(9)
        picture = logo.add_run().add_picture(str(LOGO_PATH), height=Inches(1.35))
        set_inline_picture_alt(picture, "Biểu trưng Trường Điện - Điện tử, Đại học Bách khoa Hà Nội")
    centered("BÁO CÁO BÀI TẬP LỚN HỌC PHẦN KỸ THUẬT PHẦN MỀM ỨNG DỤNG", 17, True, 10)
    centered("Đề tài: Xây dựng hệ thống Custom Keyboard Builder", 16, True, 20)

    metadata = doc.add_table(rows=8, cols=3)
    set_table_geometry(metadata, [2100, 5060, 2200], indent_dxa=40)
    remove_table_borders(metadata)
    rows = [
        ("GVHD:", "........................................................", ""),
        ("Mã lớp:", "........................................................", ""),
        ("Nhóm sinh viên thực hiện:", "", ""),
        ("1.", "........................................................", "MSSV: ................"),
        ("2.", "........................................................", "MSSV: ................"),
        ("3.", "........................................................", "MSSV: ................"),
        ("4.", "........................................................", "MSSV: ................"),
        ("5.", "........................................................", "MSSV: ................"),
    ]
    for row, values in zip(metadata.rows, rows):
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_margins(cell, top=45, bottom=45, start=40, end=40)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, size=12.5, bold=(idx == 0 and row is not metadata.rows[2]))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)
    centered("Hà Nội, 7/2026", 13, True, 0)


def add_front_matter(doc: Document, page_map: dict[str, int], manager: BookmarkManager) -> None:
    add_front_title(doc, "MỤC LỤC")
    for entry in TOC_ENTRIES:
        page = str(page_map.get(entry.anchor, "00"))
        add_nav_line(
            doc, entry.text, entry.anchor, page, entry.level,
            size=12.2 if entry.level == 1 else 11.6 if entry.level == 2 else 11,
            bold=entry.level == 1,
        )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_front_title(doc, "DANH MỤC HÌNH VẼ")
    for figure in FIGURES:
        add_nav_line(doc, figure.label, figure.anchor, str(page_map.get(figure.anchor, "00")),
                     level=1, size=11.2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_front_title(doc, "DANH MỤC BẢNG BIỂU")
    for anchor, (number, title) in TABLES.items():
        add_nav_line(doc, f"Bảng {number}. {title}", anchor, str(page_map.get(anchor, "00")),
                     level=1, size=11.4)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_front_title(doc, "DANH MỤC KÝ HIỆU VÀ CHỮ VIẾT TẮT")
    abbreviations = [
        "CSDL: Cơ sở dữ liệu.",
        "DBML: Database Markup Language - ngôn ngữ mô tả mô hình cơ sở dữ liệu.",
        "ERD: Entity Relationship Diagram - sơ đồ thực thể quan hệ.",
        "FK: Foreign Key - khóa ngoại.",
        "GUI: Graphical User Interface - giao diện người dùng.",
        "MQTT: Message Queuing Telemetry Transport - giao thức truyền thông publish/subscribe.",
        "MVVM: Model-View-ViewModel - mẫu kiến trúc giao diện.",
        "PK: Primary Key - khóa chính.",
        "QC: Quality Control - kiểm tra chất lượng.",
        "SQL: Structured Query Language - ngôn ngữ truy vấn có cấu trúc.",
        "SSMS: SQL Server Management Studio.",
        "WPF: Windows Presentation Foundation.",
        "XAML: Extensible Application Markup Language.",
    ]
    for item in abbreviations:
        add_bullet(doc, item)


def add_opening(doc: Document, manager: BookmarkManager, page_count: int | str) -> None:
    add_heading(doc, manager, "opening", "MỞ ĐẦU", 1, page_break_before=True)
    add_heading(doc, manager, "opening_1", "1. Giới thiệu và lý do chọn đề tài", 2)
    add_body(
        doc,
        "Bàn phím cơ tùy chỉnh được hình thành từ nhiều nhóm linh kiện có quan hệ tương thích như keyboard kit, switch, keycap, stabilizer, phụ kiện và các tùy chọn mod. Khi cấu hình được quản lý bằng ghi chú rời rạc, sai lệch về công nghệ switch, kiểu mount, form factor, số lượng switch và giá tại thời điểm đặt hàng có thể phát sinh.",
    )
    add_body(
        doc,
        "Đề tài Custom Keyboard Builder được lựa chọn nhằm số hóa quy trình từ cấu hình build đến gửi yêu cầu lắp ráp, xử lý bởi Seller, kiểm tra QC và theo dõi trạng thái. Dữ liệu được tập trung trên SQL Server, còn giao diện desktop WPF hỗ trợ ba vai trò Buyer, Seller và Admin trong cùng một hệ thống.",
    )

    add_heading(doc, manager, "opening_2", "2. Mục đích và nhiệm vụ của đề tài", 2)
    add_body(
        doc,
        "Mục đích của đề tài là xây dựng ứng dụng giúp tạo và quản lý cấu hình bàn phím tùy chỉnh theo mô hình kit-based. Hệ thống hỗ trợ kiểm tra sự phù hợp cơ bản giữa các linh kiện, ghi lại giá tại thời điểm tạo cấu hình, gửi yêu cầu lắp ráp đến Seller và quản lý dữ liệu theo từng vai trò.",
    )
    tasks = [
        "Xác định các vai trò sử dụng, chức năng cần có và quy trình hoạt động chính của hệ thống.",
        "Thiết kế cơ sở dữ liệu gồm 21 bảng và 33 khóa ngoại.",
        "Xây dựng giao diện WPF theo mô hình MVVM cho đăng nhập, Buyer Dashboard, Seller Dashboard, Admin Dashboard và chat.",
        "Xây dựng các chức năng xử lý tài khoản, danh mục linh kiện, cấu hình bàn phím, yêu cầu lắp ráp, đăng ký Seller, chat, nhật ký hệ thống và kiểm tra chất lượng mô phỏng (QC).",
        "Tạo các tập lệnh cơ sở dữ liệu và kiểm tra mã nguồn trước khi hoàn thiện báo cáo.",
    ]
    for task in tasks:
        add_bullet(doc, task)

    add_heading(doc, manager, "opening_3", "3. Phân công nhiệm vụ của nhóm", 2)
    add_linked_body(
        doc,
        [
            ("Các vị trí họ tên và mã số sinh viên được để trống để hoàn thiện theo danh sách nhóm chính thức. Cách phân chia công việc đề xuất được trình bày tại ", None),
            ("Bảng 0.1", "tbl_0_1"),
            (".", None),
        ],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_0_1")
    add_data_table(
        doc,
        ["Thành viên", "Nhiệm vụ được phân công", "Sản phẩm/tiêu chí hoàn thành"],
        [
            ["Họ tên - MSSV: ........................", "Phân tích yêu cầu; xây dựng Use Case, Sequence và Activity Diagram.", "Đặc tả tác nhân, phạm vi và luồng nghiệp vụ."],
            ["Họ tên - MSSV: ........................", "Thiết kế ERD, schema SQL Server và dữ liệu seed.", "ERD 21 bảng/33 FK; DDL theo thứ tự phụ thuộc."],
            ["Họ tên - MSSV: ........................", "Triển khai WPF View, ViewModel, theme và điều hướng theo vai trò.", "Login và dashboard Buyer/Seller/Admin."],
            ["Họ tên - MSSV: ........................", "Triển khai service, repository, validation, MQTT và QC mô phỏng.", "Luồng build/request/chat/QC hoạt động theo thiết kế."],
            ["Họ tên - MSSV: ........................", "Kiểm thử, seed/verify, tổng hợp tài liệu và báo cáo.", "Build, source gate, ma trận kiểm thử và bản Word."],
        ],
        [2100, 4000, 3260],
    )

    add_heading(doc, manager, "opening_4", "4. Nội dung sơ lược của báo cáo", 2)
    add_body(
        doc,
        f"Báo cáo gồm {page_count} trang, 3 chương chính, {len(FIGURES)} hình vẽ, {len(TABLES)} bảng biểu và 25 khối mã nguồn. Trong số các hình vẽ, {len([figure for figure in FIGURES if figure.group != 'UI'])} vị trí diagram ở Chương 2 được để trống để chèn thủ công; 2 hình giao diện được minh họa ở Chương 3.",
    )
    add_bullet(doc, "Chương 1 trình bày bối cảnh, yêu cầu về dữ liệu, mục tiêu, phạm vi và yêu cầu của hệ thống.")
    add_bullet(doc, "Chương 2 trình bày công nghệ, kiến trúc và bốn loại diagram: Entity-Relationship Diagram (ERD), Use Case Diagram, Sequence Diagram và Activity Diagram.")
    add_bullet(doc, "Chương 3 trình bày 21 đoạn DDL theo từng bảng, giao diện/mã nguồn tiêu biểu, kết quả xác minh và mức độ hoàn thành.")


def add_chapter_1(doc: Document, manager: BookmarkManager) -> None:
    add_heading(doc, manager, "ch1", "CHƯƠNG 1. TỔNG QUAN HỆ THỐNG CUSTOM KEYBOARD BUILDER", 1)
    add_body(
        doc,
        "Chương này trình bày bối cảnh hình thành, mục tiêu, phạm vi và các yêu cầu chính của hệ thống. Kết quả tổng quan được sử dụng làm cơ sở cho phần phân tích thiết kế ở chương tiếp theo.",
        italic=True,
    )

    add_heading(doc, manager, "ch1_1", "1.1. Bối cảnh và vấn đề cần giải quyết", 2)
    add_body(
        doc,
        "Một bộ bàn phím custom thường yêu cầu lựa chọn nhiều linh kiện và cân nhắc quan hệ tương thích. Mô hình cũ tách case, PCB và plate thành các bảng độc lập, đồng thời duy trì compatibility rule và tồn kho riêng theo Seller; phạm vi này tạo ra số lượng quan hệ lớn và khó hoàn thành trong khuôn khổ môn học.",
    )
    add_body(
        doc,
        "Mô hình hiện tại lấy keyboard kit làm nền tảng. Kit đã bao gồm case, PCB, plate và các phần đi kèm; Buyer tiếp tục chọn switch, keycap, stabilizer, accessory và mod note. Mẫu build_items cho phép mỗi dòng tham chiếu đúng một loại sản phẩm, còn snapshot pricing giữ nguyên giá tại thời điểm cấu hình hoặc gửi request.",
    )

    add_heading(doc, manager, "ch1_data", "1.2. Yêu cầu về dữ liệu", 2)
    add_body(
        doc,
        "Dữ liệu của hệ thống cần được lưu tập trung trên SQL Server, có liên kết rõ ràng và đủ thông tin để kiểm tra lại quá trình tạo cấu hình, xử lý yêu cầu lắp ráp và kiểm tra chất lượng.",
    )
    add_linked_body(
        doc,
        [("Các nhóm dữ liệu cần quản lý được tổng hợp tại ", None), ("Bảng 1.1", "tbl_1_data"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_1_data")
    data_requirements_table = add_data_table(
        doc,
        ["STT", "Nhóm dữ liệu", "Dữ liệu cần lưu", "Yêu cầu/ràng buộc"],
        [
            [
                "1",
                "Tài khoản, phân quyền và Seller",
                "Vai trò, tên đăng nhập, email, số điện thoại, mật khẩu đã băm, trạng thái tài khoản, hồ sơ Seller và đơn đăng ký Seller.",
                "Tên đăng nhập, email và số điện thoại không được trùng. Tài khoản không hoạt động không được đăng nhập.",
            ],
            [
                "2",
                "Danh mục linh kiện",
                "Thương hiệu, layout, keyboard kit, switch, keycap, stabilizer, phụ kiện, giá và trạng thái khả dụng.",
                "Giá không được âm; số switch yêu cầu của kit phải lớn hơn 0. Dữ liệu phải hỗ trợ lọc và kiểm tra tương thích cơ bản.",
            ],
            [
                "3",
                "Cấu hình và yêu cầu lắp ráp",
                "Build, linh kiện đã chọn, mod, số lượng, giá snapshot, Seller tiếp nhận, nội dung JSON, trạng thái và thời điểm xử lý request.",
                "Mỗi build item chỉ tham chiếu một loại sản phẩm. Giá và cấu hình tại thời điểm gửi request phải được giữ nguyên; trạng thái chuyển đúng luồng nghiệp vụ.",
            ],
            [
                "4",
                "Kiểm tra chất lượng",
                "Thiết bị kiểm tra, phiên QC, kết quả từng phím, độ trễ, độ ồn và kết luận đạt/cảnh báo/không đạt.",
                "Request chỉ được hoàn thành khi phiên QC gần nhất có kết quả đạt hoặc cảnh báo.",
            ],
            [
                "5",
                "Trao đổi và nhật ký",
                "Hội thoại, tin nhắn, người gửi, thời điểm gửi, request liên quan và lịch sử thay đổi trước/sau.",
                "Chỉ hỗ trợ trao đổi Buyer-Seller và Admin-Seller. Dữ liệu phải cho phép truy vết thao tác quan trọng.",
            ],
            [
                "6",
                "Toàn vẹn dữ liệu",
                "Khóa chính, khóa ngoại, trường bắt buộc, giá trị duy nhất, điều kiện kiểm tra và chỉ mục của 21 bảng nghiệp vụ.",
                "Duy trì đúng 33 liên kết khóa ngoại, hạn chế dữ liệu trùng lặp và tránh bản ghi mồ côi.",
            ],
        ],
        [880, 1740, 3260, 3480],
    )
    for row in data_requirements_table.rows[1:]:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_body(
        doc,
        "Những yêu cầu này là cơ sở để xây dựng ERD tổng quát, ERD chi tiết và các bảng SQL được trình bày ở các chương sau.",
    )

    add_heading(doc, manager, "ch1_2", "1.3. Mục tiêu và phạm vi hệ thống", 2)
    add_linked_body(
        doc,
        [("Ba tác nhân chính và mục tiêu sử dụng được tổng hợp tại ", None), ("Bảng 1.2", "tbl_1_1"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_1_1")
    add_data_table(
        doc,
        ["Tác nhân", "Mục tiêu", "Dữ liệu/chức năng chính"],
        [
            ["Buyer", "Tạo cấu hình và đặt lắp ráp bàn phím.", "Build, linh kiện, request, Seller application, chat, QC summary."],
            ["Seller", "Tiếp nhận và xử lý request được gán.", "Request payload, state machine, QC, analytics, chat."],
            ["Admin", "Quản trị người dùng và dữ liệu nền.", "User/role, Seller, catalog, đơn Seller, audit log, chat."],
        ],
        [1500, 3100, 4760],
    )

    add_heading(doc, manager, "ch1_3", "1.4. Tác nhân và yêu cầu chức năng", 2)
    add_linked_body(
        doc,
        [("Các yêu cầu chức năng cốt lõi được liệt kê tại ", None), ("Bảng 1.3", "tbl_1_2"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_1_2")
    add_data_table(
        doc,
        ["Nhóm chức năng", "Yêu cầu", "Vai trò"],
        [
            ["Tài khoản", "Đăng ký Buyer, đăng nhập bằng username/email, xem hồ sơ, đăng xuất và chặn tài khoản inactive.", "Guest/Buyer/Seller/Admin"],
            ["Cấu hình build", "Chọn kit, switch, keycap, stabilizer, accessory, mod; validate; tính tổng và lưu build.", "Buyer"],
            ["Request", "Chọn Seller verified, tạo JSON snapshot, gửi và theo dõi request.", "Buyer/Seller"],
            ["Xử lý Seller", "Chuyển Pending - Accepted - In_progress - Completed/Cancelled; hoàn thành sau QC hợp lệ.", "Seller"],
            ["Quản trị", "Quản lý user, Seller, catalog, đơn Seller và xem audit log.", "Admin"],
            ["Chat", "Chỉ cho phép Buyer-Seller và Admin-Seller; lưu tin nhắn trong SQL Server.", "Buyer/Seller/Admin"],
            ["QC", "Tạo session, mô phỏng telemetry từng phím, đánh giá latency/noise và tổng hợp pass/warning/fail.", "Seller/Device simulator"],
        ],
        [1850, 5660, 1850],
    )

    add_heading(doc, manager, "ch1_4", "1.5. Yêu cầu phi chức năng", 2)
    add_linked_body(
        doc,
        [("Các yêu cầu chất lượng và vận hành được trình bày tại ", None), ("Bảng 1.4", "tbl_1_3"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_1_3")
    non_functional_table = add_data_table(
        doc,
        ["Mã", "Nhóm yêu cầu", "Nội dung áp dụng"],
        [
            [
                "NFR-01",
                "Hiệu năng",
                "Các thao tác đăng nhập, tải dữ liệu bảng điều khiển, lưu cấu hình và cập nhật yêu cầu phải được xử lý bất đồng bộ để giao diện không bị treo trong thời gian chờ SQL Server. Các trường thường dùng để lọc dữ liệu cần có chỉ mục phù hợp.",
            ],
            [
                "NFR-02",
                "Bảo mật và phân quyền",
                "Mật khẩu phải được băm bằng PBKDF2-SHA256 với muối ngẫu nhiên; tài khoản ngừng hoạt động phải bị từ chối đăng nhập. Quyền theo vai trò, quyền sở hữu dữ liệu và tham số truy vấn SQL phải được kiểm tra trước khi thực hiện thao tác.",
            ],
            [
                "NFR-03",
                "Khả năng sử dụng",
                "Thông tin, nút chức năng và thông báo trạng thái phải rõ ràng, nhất quán theo từng vai trò. Giao diện phải hỗ trợ tiếng Việt và tiếng Anh, ghi nhớ ngôn ngữ đã chọn và hiển thị lỗi dữ liệu ngay trong luồng thao tác.",
            ],
            [
                "NFR-04",
                "Khả dụng và tích hợp",
                "SQL Server phải là nguồn dữ liệu chính. MQTT chỉ đóng vai trò thông báo tùy chọn; khi broker không khả dụng, các luồng nghiệp vụ cốt lõi vẫn phải tiếp tục theo chế độ chỉ sử dụng cơ sở dữ liệu.",
            ],
            [
                "NFR-05",
                "Bảo trì và mở rộng",
                "Mã nguồn phải tách View, ViewModel, Service và Repository theo MVVM; các phụ thuộc được khai báo qua interface, còn theme và tài nguyên ngôn ngữ được quản lý riêng để thuận tiện sửa đổi hoặc bổ sung chức năng.",
            ],
            [
                "NFR-06",
                "Toàn vẹn và truy vết dữ liệu",
                "Khóa chính, khóa ngoại và các ràng buộc phải duy trì liên kết hợp lệ giữa 21 bảng và 33 khóa ngoại. Yêu cầu lắp ráp phải lưu JSON snapshot; thao tác quản trị quan trọng phải ghi audit log kèm người thực hiện, thời điểm và dữ liệu thay đổi.",
            ],
            [
                "NFR-07",
                "Khả năng kiểm thử và xác minh",
                "Schema và mã nguồn phải được xác minh bằng bộ kiểm tra mã nguồn, SQL verifier và schema startup guard. Trước khi phát hành phải chạy build cùng các kịch bản kiểm thử chính cho Buyer, Seller và Admin.",
            ],
        ],
        [1100, 2200, 6060],
    )
    for row in non_functional_table.rows[1:]:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, manager, "ch1_5", "1.6. Quy tắc nghiệp vụ và ranh giới đề tài", 2)
    add_linked_body(
        doc,
        [("Phạm vi được chốt để giữ thiết kế phù hợp với mục tiêu môn học, như mô tả tại ", None), ("Bảng 1.5", "tbl_1_4"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_1_4")
    add_data_table(
        doc,
        ["Nội dung trong phạm vi", "Nội dung ngoài phạm vi/giới hạn"],
        [
            ["Cấu hình kit-based; snapshot giá; request theo Seller.", "Không chọn case/PCB/plate rời; không quản lý tồn kho riêng của Seller."],
            ["QC mô phỏng bằng C# và lưu kết quả từng phím.", "Không sử dụng ESP32, Arduino hoặc phần cứng nhúng thật."],
            ["Chat lưu SQL Server; MQTT thông báo request/status tùy chọn.", "SignalR chat realtime chưa được triển khai."],
            ["Schema/migration/verifier cho 21 PK id.", "Runtime database tại thời điểm audit chưa áp dụng migration PK-to-id."],
        ],
        [4680, 4680],
    )

    add_heading(doc, manager, "ch1_6", "1.7. Kết luận chương", 2)
    add_body(
        doc,
        "Chương 1 đã xác định được bài toán, yêu cầu về dữ liệu, tác nhân, mục tiêu, yêu cầu chức năng và ranh giới triển khai của Custom Keyboard Builder. Từ cơ sở này, Chương 2 tiếp tục mô tả công nghệ, kiến trúc, mô hình dữ liệu và các diagram dùng để chuyển yêu cầu thành thiết kế kỹ thuật.",
    )


def add_figure_group(doc: Document, manager: BookmarkManager, group: str) -> None:
    for figure in [item for item in FIGURES if item.group == group]:
        add_blank_diagram(doc, manager, figure)


def add_chapter_2(doc: Document, manager: BookmarkManager) -> None:
    add_heading(doc, manager, "ch2", "CHƯƠNG 2. PHÂN TÍCH THIẾT KẾ HỆ THỐNG", 1)
    add_body(
        doc,
        "Chương này trình bày công nghệ, kiến trúc phần mềm và bốn loại diagram được sử dụng trong báo cáo: Entity-Relationship Diagram (ERD), Use Case Diagram, Sequence Diagram và Activity Diagram. Các vùng hình diagram được để trống có caption và nguồn tương ứng để thay thế thủ công mà không làm thay đổi cấu trúc báo cáo.",
        italic=True,
    )

    add_heading(doc, manager, "ch2_1", "2.1. Công nghệ và môi trường triển khai", 2)
    add_linked_body(
        doc,
        [("Thành phần công nghệ được tổng hợp tại ", None), ("Bảng 2.1", "tbl_2_1"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_2_1")
    add_data_table(
        doc,
        ["Thành phần", "Công nghệ/phiên bản", "Vai trò"],
        [
            ["Ứng dụng desktop", "WPF trên .NET 10.0 Windows", "Xây dựng Login, Buyer/Seller/Admin Dashboard, Chat và User Menu."],
            ["Cơ sở dữ liệu", "Microsoft SQL Server Express; instance KHOADZS1VN\\SQLEXPRESS; database CustomKeyboard_Refactor", "Lưu tài khoản, catalog, build, request, QC, chat và audit."],
            ["Data access", "Microsoft.Data.SqlClient 6.1.1", "Thực hiện SQL parameterized trong repository."],
            ["Biểu đồ", "LiveChartsCore.SkiaSharpView.WPF 2.0.4", "Hiển thị KPI và analytics ở dashboard."],
            ["Realtime tùy chọn", "MQTTnet 4.3.7; broker mặc định localhost:1883", "Thông báo request/status và telemetry theo best-effort."],
            ["Bảo mật", "PBKDF2-SHA256", "Băm mật khẩu; không lưu mật khẩu rõ."],
            ["Đa ngôn ngữ", "Resource/localization nội bộ", "Chuyển đổi giao diện Việt-Anh."],
        ],
        [1900, 3650, 3810],
    )
    add_body(
        doc,
        "Không sử dụng ESP32 hoặc phần cứng nhúng trong phiên bản hiện tại. Trạm QC được mô phỏng bằng DeviceSimulator; dữ liệu có thể được phát qua MQTT nhưng SQL Server vẫn là nguồn dữ liệu chính.",
    )

    add_heading(doc, manager, "ch2_2", "2.2. Kiến trúc phần mềm", 2)
    add_body(
        doc,
        "Ứng dụng được tổ chức theo MVVM và các lớp service/repository. MainWindow đóng vai trò composition root, khởi tạo kết nối SQL Server, repository, service, MQTT, DeviceSimulator và MainShellViewModel; dashboard được điều hướng theo role sau khi đăng nhập.",
    )
    add_linked_body(
        doc,
        [("Trách nhiệm của từng lớp được nêu tại ", None), ("Bảng 2.2", "tbl_2_2"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_2_2")
    add_data_table(
        doc,
        ["Lớp", "Thành phần", "Trách nhiệm"],
        [
            ["View", "Views/*.xaml", "Khai báo bố cục, binding, command, DataGrid, chart và trạng thái hiển thị."],
            ["ViewModel", "ViewModels/*ViewModel.cs", "Quản lý state giao diện, command bất đồng bộ và điều phối service."],
            ["Service", "Services/*.cs", "Thực thi quy tắc nghiệp vụ, validation, state machine và phân quyền."],
            ["Repository", "Repositories/SqlServer/*.cs", "Đọc/ghi SQL Server bằng Microsoft.Data.SqlClient và ánh xạ model."],
            ["Data/Realtime", "Data/SqlServer, Realtime", "Tạo connection, kiểm tra schema, publish/subscribe MQTT best-effort."],
        ],
        [1400, 3160, 4800],
    )

    add_heading(doc, manager, "ch2_3", "2.3. Entity-Relationship Diagram (ERD)", 2)
    add_body(
        doc,
        "ERD chuẩn là file DBML trong Documents_Refactor. Mô hình gồm 21 bảng nghiệp vụ, 21 khóa chính đều tên id và 33 khóa ngoại; schema_migrations là bảng hạ tầng nên không được tính vào ERD nghiệp vụ.",
    )
    add_linked_body(
        doc,
        [("Các nhóm bảng được tổng hợp tại ", None), ("Bảng 2.3", "tbl_2_3"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_2_3")
    add_data_table(
        doc,
        ["Nhóm", "Bảng", "Mục đích"],
        [
            ["Tài khoản", "roles, users, seller_profiles, seller_applications", "Role, tài khoản, hồ sơ Seller và quy trình nâng cấp Buyer thành Seller."],
            ["Catalog", "brands, layouts, keyboard_kits, switches, keycap_sets, stabilizers, accessories", "Danh mục kit-based và thông tin tương thích cơ bản."],
            ["Build/Request", "builds, build_items, build_mods, build_requests", "Cấu hình, giá snapshot, mod note và yêu cầu gửi Seller."],
            ["Device/QC", "devices, device_test_sessions, device_key_test_results", "Trạm QC, session và kết quả từng phím."],
            ["Hỗ trợ", "audit_log, chat_conversations, chat_messages", "Truy vết và trao đổi theo role."],
        ],
        [1500, 3960, 3900],
    )
    add_figure_group(doc, manager, "ERD")

    add_heading(doc, manager, "ch2_4", "2.4. Use Case Diagram", 2)
    add_body(
        doc,
        "Use Case Diagram xác định ranh giới chức năng theo ba tác nhân Buyer, Seller và Admin. Không tồn tại Use Case UC-04 Device/QC độc lập trong tài liệu nguồn; DeviceSimulator là thành phần hỗ trợ luồng Seller QC.",
    )
    add_figure_group(doc, manager, "Use Case")

    add_heading(doc, manager, "ch2_5", "2.5. Sequence Diagram", 2)
    add_body(
        doc,
        "Sequence Diagram mô tả thứ tự thông điệp giữa người dùng, View/ViewModel, Service, Repository và SQL Server. Sáu luồng lõi bao phủ tài khoản, build request, Seller/QC, quản trị, duyệt đơn Seller và chat.",
    )
    add_figure_group(doc, manager, "Sequence")

    add_heading(doc, manager, "ch2_6", "2.6. Activity Diagram", 2)
    add_body(
        doc,
        "Activity Diagram nhấn mạnh các quyết định nghiệp vụ và nhánh xử lý nhìn thấy trong quá trình sử dụng. Sáu sơ đồ bao phủ tài khoản, Buyer, Seller, Admin, chat và Device/QC.",
    )
    add_figure_group(doc, manager, "Activity")

    add_heading(doc, manager, "ch2_7", "2.7. Đối chiếu diagram với thành phần triển khai", 2)
    add_linked_body(
        doc,
        [("Nguồn của bốn loại diagram được đối chiếu tại ", None), ("Bảng 2.4", "tbl_2_4"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_2_4")
    add_data_table(
        doc,
        ["Loại", "Artifact nguồn", "Phạm vi phản ánh"],
        [
            ["ERD", "Custom_Keyboard_ERD_Realistic_Kit_Shop_Proposal.dbml", "21 bảng, 33 FK, key/constraint/index chính."],
            ["Use Case", "Custom_Keyboard_Use_Cases_Refactor.md; Draw.io 16-18", "Tác nhân và chức năng theo role."],
            ["Sequence", "Custom_Keyboard_Sequence_Diagrams_6_Core.md", "Tương tác theo thời gian của sáu luồng lõi."],
            ["Activity", "Custom_Keyboard_Activity_Diagrams_Refactor.md", "Nhánh quyết định và quy trình thao tác."],
        ],
        [1400, 4300, 3660],
    )

    add_heading(doc, manager, "ch2_8", "2.8. Kết luận chương", 2)
    add_body(
        doc,
        "Chương 2 đã xác định công nghệ, kiến trúc, mô hình dữ liệu và bốn loại diagram được sử dụng trong báo cáo. Các thiết kế này tạo nền tảng cho việc triển khai schema SQL Server, giao diện WPF và quy tắc nghiệp vụ được trình bày ở Chương 3.",
    )


LOGIN_XAML = r'''<TextBlock Text="{loc:Tr Login_EmailOrUsername}"
           Style="{StaticResource FieldLabel}"/>
<TextBox AutomationProperties.AutomationId="LoginEmailOrUsernameTextBox"
         Text="{Binding EmailOrUsername, UpdateSourceTrigger=PropertyChanged}"/>

<TextBlock Text="{loc:Tr Common_Password}"
           Style="{StaticResource FieldLabel}"/>
<PasswordBox AutomationProperties.AutomationId="LoginPasswordBox"
             behaviors:PasswordBoxBinding.Attach="True"
             behaviors:PasswordBoxBinding.BoundPassword="{Binding Password, Mode=TwoWay}"/>

<TextBlock Text="{Binding ErrorMessage}"
           Foreground="{StaticResource DangerBrush}"/>
<Button Content="{loc:Tr Login_SignIn}"
        Command="{Binding LoginCommand}"
        Style="{StaticResource PrimaryButton}"/>'''


LOGIN_VIEWMODEL = r'''private async Task LoginAsync()
{
    ErrorMessage = string.Empty;
    StatusMessage = string.Empty;

    var result = await _accountService.LoginAsync(
        EmailOrUsername,
        Password);

    if (result.Succeeded && result.User is not null)
    {
        Password = string.Empty;
        _loginSucceeded(result.User);
        return;
    }

    ErrorMessage = result.Message;
}'''


BUYER_XAML = r'''<TextBlock Text="{loc:Tr Buyer_KitLabel}"
           Style="{StaticResource FieldLabel}"/>
<ComboBox ItemsSource="{Binding Kits}"
          SelectedItem="{Binding SelectedKit}"
          DisplayMemberPath="KitName"/>

<TextBlock Text="{loc:Tr Buyer_SwitchLabel}"
           Style="{StaticResource FieldLabel}"/>
<ComboBox ItemsSource="{Binding CompatibleSwitches}"
          SelectedItem="{Binding SelectedSwitch}"
          DisplayMemberPath="SwitchName"/>

<TextBlock Text="{Binding TotalPreview,
                 StringFormat={}{0:N2} USD}"/>
<Button Content="{loc:Tr Buyer_SaveBuild}"
        Command="{Binding SaveBuildCommand}"/>
<Button Content="{loc:Tr Buyer_SendRequest}"
        Command="{Binding SendRequestCommand}"/>'''


BUYER_VIEWMODEL = r'''private async Task SaveBuildAsync()
{
    var build = CreateBuildFromCurrentSelection();
    var saved = await _buildService.SaveBuildAsync(build);
    _editingBuildId = saved.BuildId;
    await RefreshBuildsAsync();
    StatusMessage = TrFormat(
        "Buyer_BuildSaved",
        saved.Name,
        saved.TotalCostSnapshot);
}

private async Task SendRequestAsync()
{
    if (SelectedBuild is null)
        throw new InvalidOperationException(
            Tr("Buyer_SelectSavedBuildFirst"));
    if (SelectedSeller is null)
        throw new InvalidOperationException(
            Tr("Buyer_SelectSellerFirst"));

    await _requestService.SendRequestAsync(
        SelectedBuild.BuildId,
        CurrentUser.UserId,
        SelectedSeller.UserId,
        RequestNote);
    await RefreshRequestsAsync();
}'''


def add_sql_group(doc: Document, manager: BookmarkManager, sql_blocks: dict[str, str],
                  anchor: str, heading: str, table_names: list[str], start_number: int) -> int:
    add_heading(doc, manager, anchor, heading, 3)
    code_number = start_number
    for table_name in table_names:
        add_code_table(
            doc,
            f"Mã SQL 3.{code_number} - Tạo bảng {table_name}",
            sql_blocks[table_name],
            "sql",
        )
        code_number += 1
    return code_number


def add_chapter_3(doc: Document, manager: BookmarkManager) -> None:
    add_heading(doc, manager, "ch3", "CHƯƠNG 3. TRIỂN KHAI VÀ KIỂM THỬ", 1)
    add_body(
        doc,
        "Chương này trình bày DDL của từng bảng theo ERD, giao diện và mã nguồn tiêu biểu, kết quả build/xác minh cùng đánh giá mức độ hoàn thành. Tất cả đoạn code được đưa dưới dạng văn bản trong bảng hai hàng, không sử dụng ảnh chụp code.",
        italic=True,
    )

    add_heading(doc, manager, "ch3_1", "3.1. Triển khai cơ sở dữ liệu SQL Server", 2)
    add_body(
        doc,
        "CreateSchema_Refactor.sql được dùng cho môi trường clean/disposable và có hành vi xóa, tạo lại 21 bảng. Database có dữ liệu phải dùng chuỗi preflight - migration PK-to-id - postflight trên bản restore rehearsal; không được chạy clean schema trực tiếp.",
    )
    add_body(
        doc,
        "Tại thời điểm audit ngày 15/07/2026, source code, DBML, schema và repository đã thống nhất khóa chính id, nhưng runtime database vẫn dùng 21 tên PK legacy. Vì vậy migration được xem là đã chuẩn bị nhưng chưa được xác nhận triển khai trên runtime.",
    )
    add_linked_body(
        doc,
        [("Thứ tự phụ thuộc của các nhóm bảng được tóm tắt tại ", None), ("Bảng 3.1", "tbl_3_1"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_3_1")
    add_data_table(
        doc,
        ["Thứ tự", "Nhóm", "Bảng", "Số bảng"],
        [
            ["1", "Tài khoản", "roles, users, seller_profiles, seller_applications", "4"],
            ["2", "Catalog", "brands, layouts, keyboard_kits, switches, keycap_sets, stabilizers, accessories", "7"],
            ["3", "Build/Request", "builds, build_items, build_mods, build_requests", "4"],
            ["4", "Device/QC", "devices, device_test_sessions, device_key_test_results", "3"],
            ["5", "Audit/Chat", "audit_log, chat_conversations, chat_messages", "3"],
        ],
        [900, 1650, 5910, 900],
    )

    sql_blocks = extract_sql_blocks()
    number = 1
    number = add_sql_group(doc, manager, sql_blocks, "ch3_1_1", "3.1.1. Nhóm tài khoản và phân quyền", TABLE_ORDER[0:4], number)
    number = add_sql_group(doc, manager, sql_blocks, "ch3_1_2", "3.1.2. Nhóm danh mục linh kiện", TABLE_ORDER[4:11], number)
    number = add_sql_group(doc, manager, sql_blocks, "ch3_1_3", "3.1.3. Nhóm build và request", TABLE_ORDER[11:15], number)
    number = add_sql_group(doc, manager, sql_blocks, "ch3_1_4", "3.1.4. Nhóm thiết bị và kiểm tra QC", TABLE_ORDER[15:18], number)
    add_sql_group(doc, manager, sql_blocks, "ch3_1_5", "3.1.5. Nhóm audit và chat", TABLE_ORDER[18:21], number)

    add_heading(doc, manager, "ch3_2", "3.2. Triển khai giao diện tiêu biểu", 2)
    add_linked_body(
        doc,
        [("Mối liên hệ giữa giao diện và mã nguồn được tổng hợp tại ", None), ("Bảng 3.2", "tbl_3_2"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_3_2")
    add_data_table(
        doc,
        ["Giao diện", "View", "ViewModel/Service", "Nội dung minh họa"],
        [
            ["Đăng nhập", "Views/LoginView.xaml", "LoginViewModel; AccountService", "Binding username/password, LoginCommand, thông báo lỗi và điều hướng theo role."],
            ["Buyer Dashboard", "Views/BuyerDashboardView.xaml", "BuyerDashboardViewModel; BuildService; RequestService", "Chọn kit/linh kiện, tổng giá, validation, lưu build và gửi request."],
        ],
        [1500, 2260, 2700, 2900],
    )
    login_figure = next(item for item in FIGURES if item.anchor == "fig_3_1")
    buyer_figure = next(item for item in FIGURES if item.anchor == "fig_3_2")
    add_picture_figure(doc, manager, login_figure, LOGIN_IMAGE)
    add_code_table(doc, "Mã nguồn 3.22 - Form đăng nhập bằng XAML", LOGIN_XAML, "xaml")
    add_code_table(doc, "Mã nguồn 3.23 - ViewModel xử lý đăng nhập", LOGIN_VIEWMODEL, "csharp")
    add_picture_figure(doc, manager, buyer_figure, BUYER_IMAGE)
    add_code_table(doc, "Mã nguồn 3.24 - Khu vực cấu hình build bằng XAML", BUYER_XAML, "xaml")
    add_code_table(doc, "Mã nguồn 3.25 - Lưu build và gửi request trong ViewModel", BUYER_VIEWMODEL, "csharp")

    add_heading(doc, manager, "ch3_3", "3.3. Kiểm thử và xác minh", 2)
    add_linked_body(
        doc,
        [
            ("Việc xác minh được thực hiện theo phạm vi không làm thay đổi dữ liệu runtime. Kết quả hiện tại được trình bày tại ", None),
            ("Bảng 3.3", "tbl_3_3"),
            ("; các tuyên bố kiểm thử lịch sử không được dùng thay cho runner còn tồn tại trong working tree.", None),
        ],
    )
    add_table_caption(doc, manager, "tbl_3_3")
    add_data_table(
        doc,
        ["Mã", "Nội dung", "Kết quả", "Ghi chú"],
        [
            ["T01", "dotnet build", "Pass", "0 warning, 0 error; Custom_keyboard.dll được tạo cho net10.0-windows."],
            ["T02", "VerifyPkToId_Source.ps1", "Pass", "21 PK id, 33 FK, 21 rename, version/startup guard và repository scan đạt."],
            ["T03", "Database runtime postflight", "Chưa đạt", "Runtime tại thời điểm audit vẫn dùng 21 PK legacy; migration chưa chạy."],
            ["T04", "VerifyRefactor.sql trên clean-seed", "Chưa chạy trong lượt lập báo cáo", "Script có kiểm tra schema/data nhưng có thể ghi dữ liệu fixture; không dùng trên runtime legacy."],
            ["T05", "C# integration/UI automation", "Không khả dụng", "Phase6Verification và WpfUiVerification không còn trong working tree hiện tại."],
        ],
        [900, 2600, 1500, 4360],
    )

    add_heading(doc, manager, "ch3_4", "3.4. Đánh giá mức độ hoàn thành", 2)
    add_linked_body(
        doc,
        [("Mức độ hoàn thành theo từng mục tiêu được tổng hợp tại ", None), ("Bảng 3.4", "tbl_3_4"), (".", None)],
        keep_next=True,
    )
    add_table_caption(doc, manager, "tbl_3_4")
    add_data_table(
        doc,
        ["Mục tiêu", "Trạng thái", "Đánh giá"],
        [
            ["Phân tích thiết kế", "Đạt", "Có ERD, Use Case, Sequence, Activity và mapping artifact."],
            ["Giao diện và nghiệp vụ ba vai trò", "Đạt ở mức mã nguồn", "View/ViewModel/Service/Repository cho Buyer, Seller, Admin đã được triển khai."],
            ["Schema SQL Server", "Đạt ở source", "DBML/schema/migration/verifier đã đồng bộ 21 PK id và 33 FK."],
            ["Vận hành trên runtime database", "Chưa đạt", "Cần backup, restore rehearsal, migration, postflight và smoke test."],
            ["QC phần cứng", "Ngoài phạm vi", "QC hiện dùng DeviceSimulator; không có ESP32/phần cứng thật."],
            ["Chat realtime SignalR", "Chưa triển khai", "Chat đã lưu DB; SignalR là phần mở rộng tùy chọn."],
        ],
        [3000, 1800, 4560],
    )

    heading = add_heading(doc, manager, "ch3_5", "3.5. Kết luận chương", 2)
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(2)
    paragraph = add_body(
        doc,
        "Chương 3 đã trình bày 21 DDL, hai giao diện, bốn đoạn mã và kết quả xác minh. Nghiệm thu vận hành còn chờ migration runtime và smoke test.",
        size=12,
    )
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)


def add_conclusion(doc: Document, manager: BookmarkManager) -> None:
    add_heading(doc, manager, "conclusion", "KẾT LUẬN", 1)
    add_heading(doc, manager, "conclusion_1", "1. Mức độ hoàn thành nhiệm vụ", 2)
    add_body(
        doc,
        "Các nhiệm vụ phân tích, thiết kế và triển khai mã nguồn cốt lõi đã được hoàn thành phần lớn: ứng dụng WPF theo MVVM, phân quyền Buyer/Seller/Admin, cấu hình build kit-based, request state machine, Seller application, chat lưu database, audit, analytics và QC mô phỏng đều có thành phần triển khai tương ứng.",
    )
    add_body(
        doc,
        "Tuy nhiên, nhiệm vụ chưa thể được coi là hoàn thành ở mức vận hành cuối cùng. Runtime database được audit ngày 15/07/2026 vẫn dùng schema PK legacy, do đó cần hoàn tất backup, restore rehearsal, migration, postflight và smoke test trước khi nghiệm thu.",
    )

    add_heading(doc, manager, "conclusion_2", "2. Mức độ phù hợp với mục đích đề tài", 2)
    add_body(
        doc,
        "Thiết kế hiện tại phù hợp với mục đích số hóa quy trình cấu hình và đặt lắp ráp bàn phím tùy chỉnh. Mô hình kit-based giảm độ phức tạp so với việc chọn case/PCB/plate rời; snapshot pricing, validation và phân quyền hỗ trợ tính đúng đắn của dữ liệu; SQL Server tạo nền tảng lưu trữ tập trung và truy vết.",
    )

    add_heading(doc, manager, "conclusion_3", "3. Hạn chế còn tồn tại", 2)
    for item in [
        "Runtime database chưa áp dụng migration PK-to-id và chưa có bằng chứng smoke test sau migration.",
        "Phase6Verification và WpfUiVerification không còn trong working tree nên chưa thể chạy lại bộ C# integration/UI automation lịch sử.",
        "QC mới là mô phỏng phần mềm, chưa kết nối ESP32 hoặc trạm đo phần cứng thật.",
        "SignalR chat realtime chưa được triển khai; MQTT chủ yếu phục vụ thông báo request/status và telemetry tùy chọn.",
        "Năm cột logical enum gồm role_name và switch technology chưa có CHECK constraint đầy đủ trong schema hiện hành.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, manager, "conclusion_4", "4. Hướng phát triển", 2)
    for item in [
        "Diễn tập migration trên bản restore, chạy postflight, khôi phục runner integration/UI và ghi nhận biên bản smoke test.",
        "Bổ sung CHECK constraint, kiểm thử bảo mật/đầu vào/log, đồng thời hoàn thiện SignalR và cơ chế cancellation/versioning.",
        "Kết nối trạm QC phần cứng hoặc ESP32 khi phạm vi đề tài được mở rộng sang IoT.",
    ]:
        add_bullet(doc, item)


def build_document(output_path: Path, page_map: dict[str, int] | None = None,
                   page_count: int | str = "00") -> None:
    page_map = page_map or {}
    doc = Document()
    configure_styles(doc)
    manager = BookmarkManager()
    add_cover(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(body_section, cover=False)
    remove_page_border(body_section)
    set_page_footer(body_section)

    add_front_matter(doc, page_map, manager)
    add_opening(doc, manager, page_count)
    add_chapter_1(doc, manager)
    add_chapter_2(doc, manager)
    add_chapter_3(doc, manager)
    add_conclusion(doc, manager)

    doc.settings.update_fields_on_open = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def normalize_text(text: str) -> str:
    text = text.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text).strip().casefold()


def extract_page_map(pdf_path: Path) -> tuple[dict[str, int], int]:
    reader = PdfReader(str(pdf_path))
    labels: dict[str, str] = {entry.anchor: entry.text for entry in TOC_ENTRIES}
    labels.update({figure.anchor: figure.label for figure in FIGURES})
    labels.update({anchor: f"Bảng {number}. {title}" for anchor, (number, title) in TABLES.items()})
    normalized_labels = {anchor: normalize_text(label) for anchor, label in labels.items()}
    page_map: dict[str, int] = {}
    for page_number, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        for anchor, label in normalized_labels.items():
            if label and label in text:
                # Keep the last occurrence so actual headings/captions win over front-matter lists.
                page_map[anchor] = page_number
    missing = [anchor for anchor in labels if anchor not in page_map]
    if missing:
        raise RuntimeError(f"Could not locate report labels in rendered PDF: {missing}")
    return page_map, len(reader.pages)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the academic Custom Keyboard project report.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--from-pdf", type=Path, help="Use a rendered draft/final PDF to populate page numbers.")
    args = parser.parse_args()

    if args.from_pdf:
        page_map, page_count = extract_page_map(args.from_pdf)
        build_document(args.output, page_map=page_map, page_count=page_count)
        print(f"Built {args.output} from {args.from_pdf}: {page_count} pages, {len(page_map)} mapped labels.")
    else:
        build_document(args.output, page_map={}, page_count="00")
        print(f"Built draft {args.output} with placeholder page numbers.")


if __name__ == "__main__":
    main()
