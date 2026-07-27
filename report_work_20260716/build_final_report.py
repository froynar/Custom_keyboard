from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from tools import build_academic_project_report as base  # noqa: E402


DEFAULT_OUTPUT = ROOT / "output" / "Bao_cao_du_an_Custom_Keyboard_Hoan_chinh_20260716.docx"


# The report follows the narrative_proposal preset with a named academic-report
# override matching the supplied samples: A4 page, Times New Roman 13 pt,
# justified body, academic cover, restrained gray code headers, and static TOC.

DIAGRAMS = [
    base.FigureDef(
        "fig_2_1",
        "2.1",
        "Sơ đồ Use Case tổng hợp của hệ thống Custom Keyboard Builder",
        "Use Case",
        "Documents_Refactor/Custom_Keyboard_Use_Cases_Refactor.md; KTPMUD_diagram/16-18_UseCase_*.drawio",
        "Sơ đồ cần tổng hợp ba tác nhân Buyer, Seller và Admin; phạm vi QC được đặt trong luồng Seller, không tách thành tác nhân độc lập.",
        3.65,
    ),
    base.FigureDef(
        "fig_2_2",
        "2.2",
        "Sơ đồ Activity tổng hợp của các luồng nghiệp vụ chính",
        "Activity",
        "Documents_Refactor/Custom_Keyboard_Activity_Diagrams_Refactor.md",
        "Sơ đồ cần tổng hợp các nhánh hoạt động về tài khoản, Buyer, Seller, Admin, chat và kiểm tra Device/QC.",
        3.65,
    ),
    base.FigureDef(
        "fig_2_3",
        "2.3",
        "Sơ đồ Sequence tổng hợp của các luồng tương tác chính",
        "Sequence",
        "Documents_Refactor/Custom_Keyboard_Sequence_Diagrams_6_Core.md",
        "Sơ đồ cần tổng hợp trình tự giữa View/ViewModel, Service, Repository, SQL Server và thành phần realtime trong sáu luồng lõi.",
        3.65,
    ),
    base.FigureDef(
        "fig_2_4",
        "2.4",
        "Sơ đồ ERD tổng quát của hệ thống Custom Keyboard Builder",
        "ERD",
        "Documents_Refactor/Custom_Keyboard_ERD_Realistic_Kit_Shop_Proposal.dbml",
        "Sơ đồ cần thể hiện năm vùng dữ liệu: tài khoản, danh mục, build/request, Device/QC và audit/chat cùng các quan hệ chính.",
        3.65,
    ),
    base.FigureDef(
        "fig_2_5",
        "2.5",
        "Sơ đồ ERD chi tiết gồm 21 bảng và 33 khóa ngoại",
        "ERD",
        "Documents_Refactor/Custom_Keyboard_ERD_Realistic_Kit_Shop_Proposal.dbml; Database/SqlServer/CreateSchema_Refactor.sql",
        "Sơ đồ cần hiển thị cột, kiểu dữ liệu, khóa chính id, khóa ngoại, UNIQUE, CHECK và hai filtered/unique index có ý nghĩa nghiệp vụ.",
        3.85,
    ),
]


UI_FIGURES = [
    base.FigureDef(
        "fig_3_1",
        "3.1",
        "Giao diện đăng nhập của ứng dụng",
        "UI",
        "Views/LoginView.xaml",
        "Vùng hình dành cho ảnh chụp màn hình đăng nhập, thông báo trạng thái và điều hướng tạo tài khoản.",
        3.35,
    ),
    base.FigureDef(
        "fig_3_register",
        "3.2",
        "Giao diện đăng ký tài khoản Buyer",
        "UI",
        "Views/RegisterView.xaml",
        "Vùng hình dành cho ảnh chụp biểu mẫu đăng ký, kiểm tra dữ liệu đầu vào và phản hồi kết quả.",
        3.35,
    ),
    base.FigureDef(
        "fig_3_2",
        "3.3",
        "Giao diện Buyer Dashboard và khu vực cấu hình build",
        "UI",
        "Views/BuyerDashboardView.xaml",
        "Vùng hình dành cho ảnh chụp configurator kit-based, tổng giá snapshot, validation và thao tác gửi request.",
        3.35,
    ),
    base.FigureDef(
        "fig_3_seller",
        "3.4",
        "Giao diện Seller Dashboard và xử lý yêu cầu lắp ráp",
        "UI",
        "Views/SellerDashboardView.xaml",
        "Vùng hình dành cho ảnh chụp danh sách request, chuyển trạng thái và khu vực thực hiện kiểm tra QC.",
        3.35,
    ),
    base.FigureDef(
        "fig_3_admin",
        "3.5",
        "Giao diện Admin Dashboard và quản trị hệ thống",
        "UI",
        "Views/AdminDashboardView.xaml",
        "Vùng hình dành cho ảnh chụp KPI, quản lý tài khoản, Seller, catalog, đơn đăng ký và audit log.",
        3.35,
    ),
    base.FigureDef(
        "fig_3_chat",
        "3.6",
        "Giao diện trao đổi tin nhắn theo vai trò",
        "UI",
        "Views/ChatView.xaml",
        "Vùng hình dành cho ảnh chụp danh sách hội thoại và nội dung chat Buyer-Seller hoặc Admin-Seller.",
        3.35,
    ),
    base.FigureDef(
        "fig_3_user_menu",
        "3.7",
        "Giao diện hồ sơ và menu người dùng",
        "UI",
        "Views/UserMenuView.xaml",
        "Vùng hình dành cho ảnh chụp thông tin tài khoản, lựa chọn ngôn ngữ và thao tác đăng xuất.",
        3.15,
    ),
]


base.FIGURES = DIAGRAMS + UI_FIGURES


def _rebuild_toc_entries() -> None:
    rebuilt: list[base.TocEntry] = []
    inserted = False
    skipped = {"ch2_3", "ch2_4", "ch2_5", "ch2_6"}
    for entry in base.TOC_ENTRIES:
        if entry.anchor in skipped:
            continue
        if entry.anchor == "ch2_7" and not inserted:
            rebuilt.extend(
                [
                    base.TocEntry("ch2_3", 2, "2.3. Use Case Diagram"),
                    base.TocEntry("ch2_4", 2, "2.4. Activity Diagram"),
                    base.TocEntry("ch2_5", 2, "2.5. Sequence Diagram"),
                    base.TocEntry("ch2_6", 2, "2.6. Entity-Relationship Diagram (ERD)"),
                    base.TocEntry("ch2_6_1", 3, "2.6.1. Sơ đồ ERD tổng quát"),
                    base.TocEntry("ch2_6_2", 3, "2.6.2. Sơ đồ ERD chi tiết"),
                ]
            )
            inserted = True
        rebuilt.append(entry)
    base.TOC_ENTRIES = rebuilt


_rebuild_toc_entries()


def add_final_chapter_2(doc: Document, manager: base.BookmarkManager) -> None:
    base.add_heading(doc, manager, "ch2", "CHƯƠNG 2. PHÂN TÍCH THIẾT KẾ HỆ THỐNG", 1)
    base.add_body(
        doc,
        "Chương này trình bày công nghệ, kiến trúc phần mềm và năm sơ đồ thuộc bốn nhóm ký pháp. "
        "Thứ tự Use Case - Activity - Sequence - ERD tổng quát - ERD chi tiết được lựa chọn để đi từ chức năng, "
        "quy trình và tương tác đến cấu trúc dữ liệu; phần ERD đồng thời tạo cầu nối trực tiếp sang DDL ở Chương 3. "
        "Mỗi vị trí diagram được để trống, có mô tả nguồn và chú thích để chèn ảnh thủ công.",
        italic=True,
    )

    base.add_heading(doc, manager, "ch2_1", "2.1. Công nghệ và môi trường triển khai", 2)
    base.add_linked_body(
        doc,
        [("Thành phần công nghệ được tổng hợp tại ", None), ("Bảng 2.1", "tbl_2_1"), (".", None)],
        keep_next=True,
    )
    base.add_table_caption(doc, manager, "tbl_2_1")
    base.add_data_table(
        doc,
        ["Thành phần", "Công nghệ/phiên bản", "Vai trò"],
        [
            ["Ứng dụng desktop", "WPF trên .NET 10.0 Windows", "Xây dựng Login, Buyer/Seller/Admin Dashboard, Chat và User Menu."],
            ["Cơ sở dữ liệu", "Microsoft SQL Server Express; instance KHOADZS1VN\\SQLEXPRESS; database CustomKeyboard_Refactor", "Lưu tài khoản, catalog, build, request, QC, chat và audit."],
            ["Truy cập dữ liệu", "Microsoft.Data.SqlClient 6.1.1; ADO.NET", "Thực hiện câu lệnh SQL có tham số trong lớp repository; không sử dụng Entity Framework."],
            ["Biểu đồ", "LiveChartsCore.SkiaSharpView.WPF 2.0.4", "Hiển thị KPI và analytics tại dashboard."],
            ["Realtime tùy chọn", "MQTTnet 4.3.7; broker mặc định localhost:1883", "Phát thông báo request/status và telemetry theo cơ chế best-effort."],
            ["Bảo mật", "PBKDF2-SHA256", "Băm mật khẩu trước khi lưu; không lưu mật khẩu rõ."],
            ["Đa ngôn ngữ", "Resource/localization nội bộ", "Chuyển đổi giao diện Việt-Anh."],
        ],
        [1900, 3650, 3810],
    )
    base.add_body(
        doc,
        "Không sử dụng ESP32 hoặc phần cứng nhúng trong phiên bản hiện tại. Trạm QC được mô phỏng bằng "
        "DeviceSimulator; MQTT chỉ là kênh thông báo tùy chọn, còn SQL Server giữ vai trò nguồn dữ liệu chính.",
    )

    base.add_heading(doc, manager, "ch2_2", "2.2. Kiến trúc phần mềm", 2)
    base.add_body(
        doc,
        "Ứng dụng được tổ chức theo MVVM kết hợp các lớp service và repository. MainWindow đóng vai trò composition root, "
        "khởi tạo kết nối SQL Server, repository, service, MQTT, DeviceSimulator và MainShellViewModel; dashboard được "
        "điều hướng theo role sau khi đăng nhập.",
    )
    base.add_linked_body(
        doc,
        [("Trách nhiệm của từng lớp được nêu tại ", None), ("Bảng 2.2", "tbl_2_2"), (".", None)],
        keep_next=True,
    )
    base.add_table_caption(doc, manager, "tbl_2_2")
    base.add_data_table(
        doc,
        ["Lớp", "Thành phần", "Trách nhiệm"],
        [
            ["View", "Views/*.xaml", "Khai báo bố cục, binding, command, DataGrid, chart và trạng thái hiển thị."],
            ["ViewModel", "ViewModels/*ViewModel.cs", "Quản lý state giao diện, command bất đồng bộ và điều phối service."],
            ["Service", "Services/*.cs", "Thực thi quy tắc nghiệp vụ, validation, state machine và phân quyền."],
            ["Repository", "Repositories/SqlServer/*.cs", "Đọc/ghi SQL Server bằng Microsoft.Data.SqlClient và ánh xạ model."],
            ["Data/Realtime", "Data/SqlServer; Realtime", "Tạo connection, kiểm tra schema và publish/subscribe MQTT best-effort."],
        ],
        [1700, 3000, 4660],
    )

    base.add_heading(
        doc,
        manager,
        "ch2_3",
        "2.3. Use Case Diagram",
        2,
        page_break_before=True,
    )
    base.add_body(
        doc,
        "Use Case Diagram xác định ranh giới chức năng theo ba tác nhân Buyer, Seller và Admin. Nguồn thiết kế chỉ có "
        "UC-01 Buyer, UC-02 Seller và UC-03 Admin; Device/QC là thành phần hỗ trợ luồng Seller, không tạo thành UC-04 độc lập.",
    )
    base.add_blank_diagram(doc, manager, DIAGRAMS[0])

    base.add_heading(doc, manager, "ch2_4", "2.4. Activity Diagram", 2)
    base.add_body(
        doc,
        "Activity Diagram mô hình hóa các bước xử lý và nhánh quyết định của sáu nhóm hoạt động: tài khoản, tạo build/request, "
        "Seller xử lý request, Admin quản trị, chat và kiểm tra Device/QC.",
    )
    base.add_blank_diagram(doc, manager, DIAGRAMS[1])

    base.add_heading(doc, manager, "ch2_5", "2.5. Sequence Diagram", 2)
    base.add_body(
        doc,
        "Sequence Diagram mô tả thứ tự thông điệp giữa người dùng, View/ViewModel, Service, Repository, SQL Server và "
        "thành phần realtime. Sáu luồng lõi bao phủ tài khoản, build request, Seller/QC, quản trị, duyệt đơn Seller và chat.",
    )
    base.add_blank_diagram(doc, manager, DIAGRAMS[2])

    base.add_heading(doc, manager, "ch2_6", "2.6. Entity-Relationship Diagram (ERD)", 2)
    base.add_body(
        doc,
        "ERD chuẩn được mô tả bằng DBML trong Documents_Refactor. Mô hình gồm 21 bảng nghiệp vụ, 21 khóa chính đều tên id "
        "và 33 khóa ngoại; dbo.schema_migrations là bảng hạ tầng nên không được tính vào ERD nghiệp vụ.",
    )
    base.add_linked_body(
        doc,
        [("Các nhóm bảng được tổng hợp tại ", None), ("Bảng 2.3", "tbl_2_3"), (".", None)],
        keep_next=True,
    )
    base.add_table_caption(doc, manager, "tbl_2_3")
    base.add_data_table(
        doc,
        ["Nhóm", "Bảng", "Mục đích"],
        [
            ["Tài khoản", "roles, users, seller_profiles, seller_applications", "Role, tài khoản, hồ sơ Seller và quy trình nâng cấp Buyer thành Seller."],
            ["Catalog", "brands, layouts, keyboard_kits, switches, keycap_sets, stabilizers, accessories", "Danh mục kit-based và thông tin tương thích cơ bản."],
            ["Build/Request", "builds, build_items, build_mods, build_requests", "Cấu hình, giá snapshot, mod note và yêu cầu gửi Seller."],
            ["Device/QC", "devices, device_test_sessions, device_key_test_results", "Trạm QC, phiên kiểm tra và kết quả từng phím."],
            ["Hỗ trợ", "audit_log, chat_conversations, chat_messages", "Truy vết và trao đổi theo role."],
        ],
        [1500, 3960, 3900],
    )
    base.add_heading(doc, manager, "ch2_6_1", "2.6.1. Sơ đồ ERD tổng quát", 3)
    base.add_blank_diagram(doc, manager, DIAGRAMS[3])
    base.add_heading(doc, manager, "ch2_6_2", "2.6.2. Sơ đồ ERD chi tiết", 3)
    base.add_blank_diagram(doc, manager, DIAGRAMS[4])

    base.add_heading(doc, manager, "ch2_7", "2.7. Đối chiếu diagram với thành phần triển khai", 2)
    base.add_linked_body(
        doc,
        [("Nguồn của năm sơ đồ thuộc bốn nhóm ký pháp được đối chiếu tại ", None), ("Bảng 2.4", "tbl_2_4"), (".", None)],
        keep_next=True,
    )
    base.add_table_caption(doc, manager, "tbl_2_4")
    base.add_data_table(
        doc,
        ["Nhóm ký pháp", "Artifact nguồn", "Phạm vi phản ánh"],
        [
            ["Use Case", "Custom_Keyboard_Use_Cases_Refactor.md; Draw.io 16-18", "Tác nhân và chức năng theo role; UC-02 bao gồm QC."],
            ["Activity", "Custom_Keyboard_Activity_Diagrams_Refactor.md", "Nhánh quyết định của sáu nhóm hoạt động."],
            ["Sequence", "Custom_Keyboard_Sequence_Diagrams_6_Core.md", "Tương tác theo thời gian của sáu luồng lõi."],
            ["ERD", "Custom_Keyboard_ERD_Realistic_Kit_Shop_Proposal.dbml", "Một sơ đồ tổng quát và một sơ đồ chi tiết cho 21 bảng, 33 FK."],
        ],
        [1600, 4180, 3580],
    )

    base.add_heading(doc, manager, "ch2_8", "2.8. Kết luận chương", 2)
    base.add_body(
        doc,
        "Chương 2 đã xác định công nghệ, kiến trúc, luồng chức năng và mô hình dữ liệu thông qua năm sơ đồ thuộc bốn nhóm ký pháp. "
        "ERD chi tiết là cơ sở trực tiếp để trình bày 21 khối DDL SQL Server, giao diện và kết quả kiểm thử tại Chương 3.",
    )


ORIGINAL_ADD_DATA_TABLE = base.add_data_table
ORIGINAL_CONFIGURE_STYLES = base.configure_styles
ORIGINAL_ADD_BLANK_DIAGRAM = base.add_blank_diagram


def set_a4_page_geometry(section, cover: bool = False) -> None:
    """Use the Vietnamese academic-report A4 geometry."""
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    if cover:
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.62)
        section.right_margin = Inches(0.62)
    else:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)


def configure_styles_academic(doc: Document) -> None:
    ORIGINAL_CONFIGURE_STYLES(doc)
    for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        if style_name in doc.styles:
            doc.styles[style_name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_nav_line_a4(
    doc: Document,
    text: str,
    anchor: str,
    page: str,
    level: int = 1,
    size: float = 12,
    bold: bool = False,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25 * (level - 1))
    paragraph.paragraph_format.first_line_indent = Inches(-0.08 if level > 1 else 0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.12),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    base.add_internal_hyperlink(
        paragraph,
        text,
        anchor,
        size=size,
        bold=bold,
        color="000000",
    )
    run = paragraph.add_run(f"\t{page}")
    base.set_run_font(run, size=size, bold=bold)


def add_blank_diagram_clean(
    doc: Document,
    manager: base.BookmarkManager,
    figure: base.FigureDef,
) -> None:
    paragraph_count = len(doc.paragraphs)
    ORIGINAL_ADD_BLANK_DIAGRAM(doc, manager, figure)
    for paragraph in doc.paragraphs[paragraph_count:]:
        if paragraph.text.startswith("Nguồn diagram:"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_front_matter_without_blank_pages(
    doc: Document,
    page_map: dict[str, int],
    manager: base.BookmarkManager,
) -> None:
    del manager

    def add_title(text: str, new_page: bool) -> None:
        base.add_front_title(doc, text)
        doc.paragraphs[-1].paragraph_format.page_break_before = new_page

    add_title("MỤC LỤC", False)
    for entry in base.TOC_ENTRIES:
        page = str(page_map.get(entry.anchor, "00"))
        base.add_nav_line(
            doc,
            entry.text,
            entry.anchor,
            page,
            entry.level,
            size=12.2 if entry.level == 1 else 11.6 if entry.level == 2 else 11,
            bold=entry.level == 1,
        )

    add_title("DANH MỤC HÌNH VẼ", True)
    for figure in base.FIGURES:
        base.add_nav_line(
            doc,
            figure.label,
            figure.anchor,
            str(page_map.get(figure.anchor, "00")),
            level=1,
            size=11.2,
        )

    add_title("DANH MỤC BẢNG BIỂU", True)
    for anchor, (number, title) in base.TABLES.items():
        base.add_nav_line(
            doc,
            f"Bảng {number}. {title}",
            anchor,
            str(page_map.get(anchor, "00")),
            level=1,
            size=11.4,
        )

    add_title("DANH MỤC KÝ HIỆU VÀ CHỮ VIẾT TẮT", True)
    abbreviations = [
        "CSDL: Cơ sở dữ liệu.",
        "DBML: Database Markup Language - ngôn ngữ mô tả mô hình cơ sở dữ liệu.",
        "ERD: Entity Relationship Diagram - sơ đồ thực thể quan hệ.",
        "FK: Foreign Key - khóa ngoại.",
        "GUI: Graphical User Interface - giao diện người dùng.",
        "MQTT: Message Queuing Telemetry Transport - giao thức truyền thông publish/subscribe.",
        "MVVM: Model-View-ViewModel - mẫu kiến trúc giao diện.",
        "PK: Primary Key - khóa chính.",
        "QC: Quality Control - kiểm tra chất lượng.",
        "SQL: Structured Query Language - ngôn ngữ truy vấn có cấu trúc.",
        "SSMS: SQL Server Management Studio.",
        "WPF: Windows Presentation Foundation.",
        "XAML: Extensible Application Markup Language.",
    ]
    for item in abbreviations:
        base.add_bullet(doc, item)


def add_data_table_with_full_ui(doc, headers, rows, widths, header_fill="E7E6E6"):
    if headers == ["Nhóm chức năng", "Yêu cầu", "Vai trò"]:
        rows = [list(row) for row in rows]
        role_labels = {
            "Tài khoản": "Guest và người dùng có tài khoản",
            "Chat": "Buyer, Seller và Admin",
            "QC": "Seller và bộ mô phỏng",
        }
        for row in rows:
            if row[0] in role_labels:
                row[2] = role_labels[row[0]]
        widths = [1850, 5260, 2250]
    if headers and headers[0] == "Giao diện":
        rows = [
            ["Đăng nhập", "Views/LoginView.xaml", "LoginViewModel; AccountService", "Nhập thông tin, xác thực, báo lỗi và điều hướng theo role."],
            ["Đăng ký", "Views/RegisterView.xaml", "RegisterViewModel; AccountService", "Đăng ký Buyer và kiểm tra dữ liệu đầu vào."],
            ["Buyer Dashboard", "Views/\nBuyerDashboard\nView.xaml", "BuyerDashboardViewModel; BuildService; RequestService", "Cấu hình build, tính giá, validation và gửi request."],
            ["Seller Dashboard", "Views/\nSellerDashboard\nView.xaml", "SellerDashboardViewModel; RequestService; DeviceService", "Xử lý request, trạng thái và QC mô phỏng."],
            ["Admin Dashboard", "Views/\nAdminDashboard\nView.xaml", "AdminDashboardViewModel; AdminService; StatsService", "KPI, user, Seller, catalog, application và audit."],
            ["Chat", "Views/ChatView.xaml", "ChatViewModel; ChatService", "Hội thoại Buyer-Seller và Admin-Seller có kiểm tra quyền."],
            ["User Menu", "Views/\nUserMenuView.xaml", "MainShellViewModel", "Hồ sơ, ngôn ngữ và đăng xuất."],
        ]
        widths = [1350, 2500, 2800, 2710]
    return ORIGINAL_ADD_DATA_TABLE(doc, headers, rows, widths, header_fill)


def add_ui_placeholders(doc, manager, figure, image_path):
    del image_path
    if figure.anchor == "fig_3_1":
        targets = [UI_FIGURES[0], UI_FIGURES[1]]
    elif figure.anchor == "fig_3_2":
        targets = UI_FIGURES[2:]
    else:
        targets = [figure]
    for target in targets:
        base.add_blank_diagram(doc, manager, target)


base.add_chapter_2 = add_final_chapter_2
base.add_data_table = add_data_table_with_full_ui
base.add_picture_figure = add_ui_placeholders
base.configure_styles = configure_styles_academic
base.set_page_geometry = set_a4_page_geometry
base.add_front_matter = add_front_matter_without_blank_pages
base.add_nav_line = add_nav_line_a4
base.add_blank_diagram = add_blank_diagram_clean


def iter_paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def patch_summary_text(output_path: Path) -> None:
    doc = Document(output_path)
    replacements = {
        "2 hình giao diện được minh họa ở Chương 3.": (
            "7 vị trí chèn ảnh giao diện ở Chương 3 được để trống để bổ sung ảnh chụp thủ công."
        ),
        "Chương 2 trình bày công nghệ, kiến trúc và bốn loại diagram: Entity-Relationship Diagram (ERD), Use Case Diagram, Sequence Diagram và Activity Diagram.": (
            "Chương 2 trình bày công nghệ, kiến trúc và năm sơ đồ thuộc bốn nhóm ký pháp theo thứ tự Use Case, Activity, Sequence, ERD tổng quát và ERD chi tiết."
        ),
        "Chương 3 đã trình bày 21 DDL, hai giao diện, bốn đoạn mã": (
            "Chương 3 đã trình bày 21 DDL, bảy vị trí chèn ảnh giao diện và bốn đoạn mã"
        ),
    }
    for paragraph in iter_paragraphs(doc):
        current = paragraph.text
        updated = current
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != current:
            if paragraph.runs:
                paragraph.runs[0].text = updated
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(updated)
    doc.save(output_path)


def build(output_path: Path, from_pdf: Path | None) -> None:
    if from_pdf:
        page_map, page_count = base.extract_page_map(from_pdf)
        base.build_document(output_path, page_map=page_map, page_count=page_count)
        patch_summary_text(output_path)
        print(f"Built {output_path} from {from_pdf}: {page_count} pages, {len(page_map)} mapped labels.")
    else:
        base.build_document(output_path, page_map={}, page_count="00")
        patch_summary_text(output_path)
        print(f"Built draft {output_path} with placeholder page numbers.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the final academic Custom Keyboard project report.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--from-pdf", type=Path)
    args = parser.parse_args()
    build(args.output, args.from_pdf)


if __name__ == "__main__":
    main()
