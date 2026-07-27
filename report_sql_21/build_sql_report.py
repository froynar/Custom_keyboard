from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_SQL = ROOT / "Database" / "SqlServer" / "CreateSchema_Refactor.sql"
OUTPUT_DIR = ROOT / "report_sql_21"
OUTPUT_DOCX = OUTPUT_DIR / "Ma_SQL_3.1.4_21_bang_ERD_moi.docx"
OUTPUT_SQL = OUTPUT_DIR / "Ma_SQL_3.1.4_21_bang_ERD_moi.sql"

TABLES = [
    "roles",
    "users",
    "seller_profiles",
    "seller_applications",
    "brands",
    "layouts",
    "keyboard_kits",
    "switches",
    "keycap_sets",
    "stabilizers",
    "accessories",
    "builds",
    "build_items",
    "build_mods",
    "build_requests",
    "devices",
    "device_test_sessions",
    "device_key_test_results",
    "audit_log",
    "chat_conversations",
    "chat_messages",
]

SQL_KEYWORDS = {
    "ADD",
    "ALTER",
    "AND",
    "AS",
    "BETWEEN",
    "CASE",
    "CHECK",
    "CONSTRAINT",
    "CREATE",
    "DEFAULT",
    "DESC",
    "ELSE",
    "END",
    "FOREIGN",
    "IDENTITY",
    "IN",
    "INDEX",
    "IS",
    "KEY",
    "NOT",
    "NULL",
    "ON",
    "OR",
    "PRIMARY",
    "REFERENCES",
    "TABLE",
    "THEN",
    "UNIQUE",
    "WHEN",
    "WHERE",
}

SQL_TYPES = {
    "BIGINT",
    "BIT",
    "DATETIME",
    "DATETIME2",
    "DECIMAL",
    "INT",
    "MAX",
    "NVARCHAR",
    "TEXT",
    "VARCHAR",
}

SQL_FUNCTIONS = {
    "LEN",
    "LTRIM",
    "RTRIM",
    "SYSUTCDATETIME",
}

TOKEN_RE = re.compile(
    r"(--[^\n]*|'(?:''|[^'])*'|\b\d+(?:\.\d+)?\b|"
    r"\b[A-Za-z_][A-Za-z0-9_]*\b|\s+|.)"
)


def strip_inline_comment(line: str) -> str:
    in_string = False
    i = 0
    while i < len(line):
        if line[i] == "'":
            if in_string and i + 1 < len(line) and line[i + 1] == "'":
                i += 2
                continue
            in_string = not in_string
        if not in_string and line[i : i + 2] == "--":
            return line[:i].rstrip()
        i += 1
    return line.rstrip()


def extract_create_statement(sql_text: str, table_name: str) -> str:
    pattern = re.compile(
        rf"CREATE\s+TABLE\s+{re.escape(table_name)}\s*\((.*?)\);\s*GO",
        flags=re.IGNORECASE | re.DOTALL,
    )
    match = pattern.search(sql_text)
    if not match:
        raise ValueError(f"Khong tim thay CREATE TABLE {table_name}")
    return f"CREATE TABLE {table_name} (\n{match.group(1)}\n);"


def normalize_create_statement(statement: str) -> str:
    cleaned: list[str] = []
    for raw_line in statement.replace("\t", "    ").splitlines():
        line = strip_inline_comment(raw_line)
        if line.strip():
            cleaned.append(line.rstrip())

    type_pattern = re.compile(
        r"^\s{2,}([A-Za-z_][A-Za-z0-9_]*)\s+"
        r"((?:BIGINT|BIT|DATETIME2?|DECIMAL|INT|NVARCHAR|TEXT|VARCHAR)"
        r"(?:\([^)]*\))?)(.*)$",
        flags=re.IGNORECASE,
    )
    column_rows: list[tuple[int, re.Match[str]]] = []
    constraints_started = False
    for index, line in enumerate(cleaned):
        stripped = line.strip()
        if stripped.upper().startswith("CONSTRAINT "):
            constraints_started = True
        if constraints_started:
            continue
        match = type_pattern.match(line)
        if match:
            column_rows.append((index, match))

    if column_rows:
        max_name = max(len(match.group(1)) for _, match in column_rows)
        max_type = max(len(match.group(2)) for _, match in column_rows)
        for index, match in column_rows:
            remainder = match.group(3).strip()
            cleaned[index] = (
                f"    {match.group(1):<{max_name}}  "
                f"{match.group(2):<{max_type}}  {remainder}"
            ).rstrip()

    return "\n".join(cleaned)


def extract_indexes(sql_text: str) -> dict[str, list[str]]:
    start = sql_text.index("-- Indexes (FK lookups")
    end = sql_text.index("\nGO", start)
    index_region = sql_text[start:end]
    statement_re = re.compile(
        r"CREATE\s+(?:UNIQUE\s+)?INDEX\b.*?;",
        flags=re.IGNORECASE | re.DOTALL,
    )
    result = {table_name: [] for table_name in TABLES}
    for match in statement_re.finditer(index_region):
        statement = match.group(0)
        on_match = re.search(
            r"\bON\s+([A-Za-z_][A-Za-z0-9_.]*)\s*\(",
            statement,
            flags=re.IGNORECASE,
        )
        if not on_match:
            raise ValueError(f"Khong xac dinh duoc bang cua index: {statement}")
        table_name = on_match.group(1).split(".")[-1]
        if table_name in result:
            result[table_name].append(normalize_index_statement(statement))
    return result


def normalize_index_statement(statement: str) -> str:
    compact = re.sub(r"\s+", " ", statement).strip()
    pattern = re.compile(
        r"CREATE\s+(UNIQUE\s+)?INDEX\s+([A-Za-z_][A-Za-z0-9_]*)\s+"
        r"ON\s+([A-Za-z_][A-Za-z0-9_.]*)\s*\(([^)]*)\)"
        r"(?:\s+WHERE\s+(.+?))?;",
        flags=re.IGNORECASE,
    )
    match = pattern.fullmatch(compact)
    if not match:
        raise ValueError(f"Khong the dinh dang index: {statement}")

    unique = "UNIQUE " if match.group(1) else ""
    columns = re.sub(r"\s*,\s*", ", ", match.group(4).strip())
    lines = [
        f"CREATE {unique}INDEX {match.group(2)}",
        f"    ON {match.group(3)}({columns})",
    ]
    if match.group(5):
        lines.append(f"    WHERE {match.group(5).strip()}")
    lines[-1] += ";"
    return "\n".join(lines)


def build_blocks(sql_text: str) -> list[dict[str, str]]:
    index_map = extract_indexes(sql_text)
    blocks: list[dict[str, str]] = []
    for number, table_name in enumerate(TABLES, start=1):
        create_sql = normalize_create_statement(
            extract_create_statement(sql_text, table_name)
        )
        parts = [create_sql]
        if index_map[table_name]:
            parts.extend(index_map[table_name])
        blocks.append(
            {
                "number": str(number),
                "table": table_name,
                "title": f"Mã SQL 3.1.4.{number} - Tạo bảng {table_name}",
                "sql": "\n\n".join(parts),
            }
        )
    return blocks


def set_repeat_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    no_proof = OxmlElement("w:noProof")
    rpr.append(no_proof)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, width_dxa: int) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    grid_col = OxmlElement("w:gridCol")
    grid_col.set(qn("w:w"), str(width_dxa))
    grid.append(grid_col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "8A96A3", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_colored_sql(paragraph, sql_text: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(8.2)
    paragraph.paragraph_format.keep_together = True

    lines = sql_text.splitlines()
    for line_index, line in enumerate(lines):
        pending_space = ""
        for token in TOKEN_RE.findall(line):
            if token.isspace():
                pending_space += token
                continue
            run = paragraph.add_run(pending_space + token)
            pending_space = ""
            set_repeat_font(run, "Consolas", 7)
            upper = token.upper()
            if token.startswith("--"):
                run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
            elif token.startswith("'"):
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            elif upper in SQL_TYPES:
                run.font.color.rgb = RGBColor(0x00, 0x7F, 0xA8)
            elif upper in SQL_KEYWORDS:
                run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
            elif upper in SQL_FUNCTIONS:
                run.font.color.rgb = RGBColor(0x79, 0x57, 0x00)
            elif re.fullmatch(r"\d+(?:\.\d+)?", token):
                run.font.color.rgb = RGBColor(0x70, 0x00, 0x70)
            else:
                run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        if pending_space:
            run = paragraph.add_run(pending_space)
            set_repeat_font(run, "Consolas", 7)
        if line_index != len(lines) - 1:
            paragraph.add_run().add_break(WD_BREAK.LINE)


def configure_document(doc: Document) -> int:
    section = doc.sections[0]
    section.start_type = WD_SECTION.CONTINUOUS
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    doc.core_properties.title = "Mã SQL 3.1.4 - 21 bảng theo ERD mới"
    doc.core_properties.subject = "SQL Server DDL cho 21 bảng nghiệp vụ"
    doc.core_properties.author = "Custom Keyboard Builder"
    doc.core_properties.keywords = "SQL Server, ERD, 21 tables, report"

    # python-docx Length values are EMU; Word table geometry expects DXA/twips.
    width_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    return width_twips


def add_sql_block(doc: Document, block: dict[str, str], width_dxa: int) -> None:
    table = doc.add_table(rows=2, cols=1)
    set_table_geometry(table, width_dxa)
    set_table_borders(table)

    header_cell = table.cell(0, 0)
    header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(header_cell, "D9D9D9")
    set_cell_margins(header_cell, top=90, start=160, bottom=90, end=160)
    header_paragraph = header_cell.paragraphs[0]
    header_paragraph.paragraph_format.space_before = Pt(0)
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    header_paragraph.paragraph_format.keep_with_next = True
    title_run = header_paragraph.add_run(block["title"])
    set_repeat_font(title_run, "Times New Roman", 12, bold=True)
    title_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    code_cell = table.cell(1, 0)
    code_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(code_cell, "FFFFFF")
    set_cell_margins(code_cell, top=100, start=180, bottom=110, end=180)
    code_paragraph = code_cell.paragraphs[0]
    add_colored_sql(code_paragraph, block["sql"])

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(8)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(2)


def write_raw_sql(blocks: list[dict[str, str]]) -> None:
    sections: list[str] = [
        "-- Custom Keyboard Builder",
        "-- Mã SQL 3.1.4: 21 bảng nghiệp vụ theo ERD mới",
        "-- Thứ tự đã bảo đảm phụ thuộc khóa ngoại.",
        "",
    ]
    for block in blocks:
        sections.append(f"-- {block['title']}")
        sections.append(block["sql"])
        sections.append("")
    OUTPUT_SQL.write_text("\n".join(sections).rstrip() + "\n", encoding="utf-8-sig")


def write_docx(blocks: list[dict[str, str]]) -> None:
    doc = Document()
    width_dxa = configure_document(doc)

    page_budget_pt = 690.0
    used_pt = 0.0
    for index, block in enumerate(blocks):
        line_count = len(block["sql"].splitlines())
        estimated_height = 28.0 + line_count * 8.2 + 12.0
        if index > 0 and used_pt + estimated_height > page_budget_pt:
            doc.add_page_break()
            used_pt = 0.0
        add_sql_block(doc, block, width_dxa)
        used_pt += estimated_height

    doc.save(OUTPUT_DOCX)


def structural_audit(blocks: list[dict[str, str]]) -> None:
    if len(blocks) != 21:
        raise AssertionError(f"Can co 21 khoi SQL, hien co {len(blocks)}")
    if [block["table"] for block in blocks] != TABLES:
        raise AssertionError("Thu tu bang khong khop ERD")
    for block in blocks:
        expected = f"CREATE TABLE {block['table']} ("
        if expected not in block["sql"]:
            raise AssertionError(f"Thieu {expected}")
    if any("schema_migrations" in block["sql"] for block in blocks):
        raise AssertionError("schema_migrations khong thuoc 21 bang nghiep vu")


def main() -> None:
    sql_text = SOURCE_SQL.read_text(encoding="utf-8-sig")
    blocks = build_blocks(sql_text)
    structural_audit(blocks)
    write_raw_sql(blocks)
    write_docx(blocks)
    print(f"Da tao: {OUTPUT_DOCX}")
    print(f"Da tao: {OUTPUT_SQL}")
    print(f"So khoi SQL: {len(blocks)}")
    for block in blocks:
        print(
            f"{int(block['number']):02d}. {block['table']}: "
            f"{len(block['sql'].splitlines())} dong"
        )


if __name__ == "__main__":
    main()
