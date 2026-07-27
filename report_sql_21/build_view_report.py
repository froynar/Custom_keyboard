from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_SQL = (
    ROOT / "Database" / "SqlServer" / "ReadViews_Simplified_20260727.sql"
)
OUTPUT_DIR = ROOT / "report_sql_21"
OUTPUT_DOCX = OUTPUT_DIR / "Ma_SQL_3.1.6_4_view_ERD_moi.docx"
OUTPUT_SQL = OUTPUT_DIR / "Ma_SQL_3.1.6_4_view_ERD_moi.sql"

VIEWS = ["Last_QC", "Catalog_Comps", "Build_items", "Req_view"]


def load_table_builder():
    builder_path = OUTPUT_DIR / "build_sql_report.py"
    spec = importlib.util.spec_from_file_location("sql_report_builder", builder_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Khong the nap builder: {builder_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BASE = load_table_builder()
BASE.SQL_KEYWORDS.update(
    {
        "ALL",
        "APPLY",
        "ASC",
        "BY",
        "CONVERT",
        "CREATE",
        "CROSS",
        "FROM",
        "FULL",
        "GROUP",
        "HAVING",
        "INNER",
        "JOIN",
        "LEFT",
        "ORDER",
        "OUTER",
        "OVER",
        "PARTITION",
        "RIGHT",
        "SELECT",
        "UNION",
        "VIEW",
        "WITH",
    }
)
BASE.SQL_FUNCTIONS.update(
    {
        "AVG",
        "COALESCE",
        "COUNT",
        "MAX",
        "NULLIF",
        "ROW_NUMBER",
        "SUM",
    }
)


def extract_view(sql_text: str, view_name: str) -> str:
    pattern = re.compile(
        rf"CREATE\s+VIEW\s+views\.{re.escape(view_name)}\s+"
        rf"AS\s+(.*?)\s*;\s*GO",
        flags=re.IGNORECASE | re.DOTALL,
    )
    match = pattern.search(sql_text)
    if not match:
        raise ValueError(f"Khong tim thay view {view_name}")

    body_lines = [line.rstrip() for line in match.group(1).splitlines()]
    while body_lines and not body_lines[0].strip():
        body_lines.pop(0)
    while body_lines and not body_lines[-1].strip():
        body_lines.pop()

    return (
        f"CREATE VIEW views.{view_name}\n"
        f"AS\n"
        f"{chr(10).join(body_lines)};\n"
        f"GO"
    )


def build_blocks(sql_text: str) -> list[dict[str, str]]:
    blocks: list[dict[str, str]] = []
    for number, view_name in enumerate(VIEWS, start=1):
        blocks.append(
            {
                "number": str(number),
                "view": view_name,
                "title": f"Mã SQL 3.1.6.{number} - Tạo view {view_name}",
                "sql": extract_view(sql_text, view_name),
            }
        )
    return blocks


def set_cell_borders(cell, **edges: dict[str, str]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge_name, settings in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key, value in settings.items():
            edge.set(qn(f"w:{key}"), value)


def visible_border() -> dict[str, str]:
    return {
        "val": "single",
        "sz": "4",
        "space": "0",
        "color": "8A96A3",
    }


def no_border() -> dict[str, str]:
    return {"val": "nil"}


def add_view_block(doc: Document, block: dict[str, str], width_dxa: int) -> None:
    # The third row is an intentionally borderless spacer. Because it belongs
    # to the copied table, Word cannot silently remove the gap by merging two
    # adjacent pasted tables.
    table = doc.add_table(rows=3, cols=1)
    BASE.set_table_geometry(table, width_dxa)

    header_cell = table.cell(0, 0)
    header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    BASE.set_cell_shading(header_cell, "D9D9D9")
    BASE.set_cell_margins(header_cell, top=90, start=160, bottom=90, end=160)
    header_paragraph = header_cell.paragraphs[0]
    header_paragraph.paragraph_format.space_before = Pt(0)
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_paragraph.paragraph_format.keep_with_next = True
    title_run = header_paragraph.add_run(block["title"])
    BASE.set_repeat_font(title_run, "Times New Roman", 12, bold=True)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    set_cell_borders(
        header_cell,
        top=visible_border(),
        start=visible_border(),
        bottom=visible_border(),
        end=visible_border(),
    )

    code_cell = table.cell(1, 0)
    code_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    BASE.set_cell_shading(code_cell, "FFFFFF")
    BASE.set_cell_margins(code_cell, top=100, start=180, bottom=110, end=180)
    code_paragraph = code_cell.paragraphs[0]
    BASE.add_colored_sql(code_paragraph, block["sql"])
    set_cell_borders(
        code_cell,
        top=visible_border(),
        start=visible_border(),
        bottom=visible_border(),
        end=visible_border(),
    )

    spacer_cell = table.cell(2, 0)
    spacer_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    BASE.set_cell_shading(spacer_cell, "FFFFFF")
    BASE.set_cell_margins(spacer_cell, top=0, start=0, bottom=0, end=0)
    spacer_paragraph = spacer_cell.paragraphs[0]
    spacer_paragraph.paragraph_format.space_before = Pt(0)
    spacer_paragraph.paragraph_format.space_after = Pt(0)
    spacer_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer_paragraph.paragraph_format.line_spacing = Pt(8)
    spacer_run = spacer_paragraph.add_run("\u00A0")
    BASE.set_repeat_font(spacer_run, "Calibri", 1)
    spacer_run.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_borders(
        spacer_cell,
        top=no_border(),
        start=no_border(),
        bottom=no_border(),
        end=no_border(),
    )


def write_docx(blocks: list[dict[str, str]]) -> None:
    doc = Document()
    width_dxa = BASE.configure_document(doc)
    doc.core_properties.title = "Mã SQL 3.1.6 - 4 view theo ERD mới"
    doc.core_properties.subject = "SQL Server read views cho báo cáo"
    doc.core_properties.author = "Custom Keyboard Builder"

    page_budget_pt = 690.0
    used_pt = 0.0
    for index, block in enumerate(blocks):
        line_count = len(block["sql"].splitlines())
        estimated_height = 34.0 + line_count * 8.2 + 12.0
        if index > 0 and used_pt + estimated_height > page_budget_pt:
            doc.add_page_break()
            used_pt = 0.0
        add_view_block(doc, block, width_dxa)
        used_pt += estimated_height

    doc.save(OUTPUT_DOCX)


def write_sql(blocks: list[dict[str, str]]) -> None:
    lines = [
        "-- Custom Keyboard Builder",
        "-- Mã SQL 3.1.6: 4 view theo hợp đồng đọc dữ liệu mới nhất",
        "-- Chạy sau khi 21 bảng nghiệp vụ đã được tạo.",
        "",
    ]
    for block in blocks:
        lines.append(f"-- {block['title']}")
        lines.append(block["sql"])
        lines.append("")
    OUTPUT_SQL.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8-sig")


def audit(blocks: list[dict[str, str]]) -> None:
    if len(blocks) != 4:
        raise AssertionError(f"Can co 4 view, hien co {len(blocks)}")
    if [block["view"] for block in blocks] != VIEWS:
        raise AssertionError("Thu tu view khong dung")
    for block in blocks:
        expected = f"CREATE VIEW views.{block['view']}"
        if expected not in block["sql"]:
            raise AssertionError(f"Thieu {expected}")
        if not block["sql"].endswith("GO"):
            raise AssertionError(f"View {block['view']} thieu GO")


def main() -> None:
    sql_text = SOURCE_SQL.read_text(encoding="utf-8-sig")
    blocks = build_blocks(sql_text)
    audit(blocks)
    write_sql(blocks)
    write_docx(blocks)
    print(f"Da tao: {OUTPUT_DOCX}")
    print(f"Da tao: {OUTPUT_SQL}")
    print(f"So view: {len(blocks)}")
    for block in blocks:
        print(
            f"{int(block['number']):02d}. {block['view']}: "
            f"{len(block['sql'].splitlines())} dong"
        )


if __name__ == "__main__":
    main()
